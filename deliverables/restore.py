"""Kembalikan data tiap borehole (γ, GWL, beban, keterangan) dan sesuaikan flag L/NL dengan FS laporan likuefaksi."""
import json, zipfile, os
import openpyxl
from xlsx_cellwriter import Sheet

GAM = {'BH-01': (17.8, 19.2), 'BH-02': (18.5, 18.4), 'BH-03': (16.5, 19.0), 'BH-04': (19.6, 20.2), 'BH-05': (18.3, 18.2)}  # Tabel 5.1
GWL = {'BH-01': 9, 'BH-02': 7, 'BH-03': 8, 'BH-04': 9, 'BH-05': 7}   # GWL analisis likuefaksi
LOADS = (5152.9, 1583.9, 1101.5, 0.0)
COPY = {'BH-04': ('BH-01', 25), 'BH-05': ('BH-02', 25), 'BH-03': ('BH-02', 22.5)}
ZLIM = 19.5          # sampel <= 19,5 m (lapisan 19,5-21 m) = batas 20 m laporan

liq = openpyxl.load_workbook('liqlat.xlsx', data_only=True)
FS = {}
for bh in GAM:
    ws = liq[bh]
    FS[bh] = {round(ws.cell(r, 2).value, 2): ws.cell(r, 7).value for r in range(90, 104)
              if isinstance(ws.cell(r, 2).value, (int, float))}
spt = json.load(open('spt.json'))

os.makedirs('drv2', exist_ok=True)
log = {}
for b in ['01', '02', '03', '04', '05']:
    bh = f'BH-{b}'
    src = f'drv/PBCP-{bh}_rev3.xlsx'
    z = zipfile.ZipFile(src)
    I, T = Sheet(z.read('xl/worksheets/sheet1.xml')), Sheet(z.read('xl/worksheets/sheet2.xml'))
    wt = openpyxl.load_workbook(src)['Data Tanah']
    ch = []
    # ---- Input ----
    I.set('C14', GWL[bh]); I.set('E14', 'GWL analisis likuefaksi (Report BP-00-GEO-RPT-003)')
    I.set('C16', GAM[bh][0]); I.set('C17', GAM[bh][1])
    I.set('E16', 'Report Tabel 5.1'); I.set('E17', 'Report Tabel 5.1')
    I.set('E26', 'sesuai perhitungan sebelumnya'); I.set('E27', 'sesuai perhitungan sebelumnya')
    I.set('E45', 'DD_max tidak difaktorkan (beban faktor 1,0)')
    if bh != 'BH-03':
        for r, v in zip((48, 49, 50, 51), LOADS):
            I.set(f'C{r}', v)
        I.set('E48', 'GROUP WPD 2 cellar 4x1, ASD2(max) - beban lama, update dgn GROUP beban revisi')
        I.set('E49', 'GROUP WPD 2 cellar 4x1, LRFD3 - beban lama, update dgn GROUP beban revisi')
        I.set('E50', 'GROUP WPD 2 cellar 4x1, ASD1 (DL) - beban lama, update dgn GROUP beban revisi')
        I.set('E51', 'GROUP: tidak ada tiang tarik pada kombinasi yang ditinjau')
    else:
        I.set('E48', 'beban uji pengguna - Logyard belum ada hasil GROUP')
    # ---- Data Tanah ----
    own = {round(t, 2): (s, n) for t, h, s, n in spt[bh]}
    for r in range(7, 48):
        top = round(float(wt.cell(r, 2).value), 2)
        if bh == 'BH-03' and top in own and top > 0:
            s, n = own[top]
            if wt.cell(r, 7).value != n or wt.cell(r, 6).value != s:
                ch.append(f'{top} m: {wt.cell(r, 6).value} N{wt.cell(r, 7).value} -> {s} N{n}')
            T.set(f'F{r}', s); T.set(f'G{r}', n)
        if bh == 'BH-05' and top in (34.5, 36.0):   # kembalikan N salinan BH-02
            T.set(f'G{r}', {34.5: 53, 36.0: 54}[top])
        fs = FS[bh].get(top)
        if bh in COPY and top >= COPY[bh][1]:
            fs = None
        flag = 'L' if (fs is not None and fs < 1.0 and top <= ZLIM) else 'NL'
        old = wt.cell(r, 8).value
        if old != flag:
            ch.append(f'{top} m: {old} -> {flag} (FS {fs if fs is None else round(fs, 2)})')
        T.set(f'H{r}', flag)
        if bh in COPY and top >= COPY[bh][1]:
            T.set(f'O{r}', f'Data salinan {COPY[bh][0]} (bor hanya s/d {COPY[bh][1]} m)')
    log[bh] = ch
    dst = f'drv2/PBCP-{bh}_rev3.xlsx'
    edited = {'xl/worksheets/sheet1.xml': I.bytes(), 'xl/worksheets/sheet2.xml': T.bytes()}
    with zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED) as zo:
        for it in z.infolist():
            zo.writestr(it, edited.get(it.filename) or z.read(it.filename))
for k, v in log.items():
    print(k, v)
