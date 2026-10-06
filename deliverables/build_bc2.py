"""Workbook baru: daya dukung, penurunan segera (immediate), konsolidasi & waktu konsolidasi per borehole - PSPE Bora Pulu.
Semua hasil berupa rumus Excel; input berwarna kuning."""
import json, statistics as st
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.chart import ScatterChart, Reference, Series
import build_bc as B          # sublapisan, klasifikasi, data lab

FONT = 'Arial'
F_T = Font(name=FONT, size=13, bold=True, color='1F3864')
F_S = Font(name=FONT, size=9, italic=True, color='595959')
F_H = Font(name=FONT, size=9, bold=True, color='FFFFFF')
F_N = Font(name=FONT, size=9)
F_B = Font(name=FONT, size=9, bold=True)
F_SEC = Font(name=FONT, size=10, bold=True, color='FFFFFF')
F_IN = Font(name=FONT, size=9, color='0000FF')
F_R = Font(name=FONT, size=10, bold=True, color='C00000')
FILL_IN = PatternFill('solid', fgColor='FFF2CC')
FILL_H = PatternFill('solid', fgColor='2F5597')
FILL_SEC = PatternFill('solid', fgColor='1F3864')
FILL_R = PatternFill('solid', fgColor='E2EFDA')
thin = Side(style='thin', color='A6A6A6')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
CEN = Alignment(horizontal='center', vertical='center', wrap_text=True)
BHS = ['BH-01', 'BH-02', 'BH-03', 'BH-04', 'BH-05']


def sec(ws, r, text, c2=12):
    ws.cell(r, 1, text).font = F_SEC
    for c in range(1, c2 + 1):
        ws.cell(r, c).fill = FILL_SEC


def put(ws, r, c, v, fmt=None, inp=False, bold=False, box=True):
    cell = ws.cell(r, c, v)
    cell.font = F_IN if inp else (F_B if bold else F_N)
    if inp:
        cell.fill = FILL_IN
    if fmt:
        cell.number_format = fmt
    if box:
        cell.border = BOX
    return cell


def line(ws, r, label, sym, formula, unit='', note='', fmt='0.00', inp=False, res=False):
    """baris perhitungan bergaya template: label | simbol | nilai | satuan | rumus/catatan"""
    ws.cell(r, 1, label).font = F_N
    ws.cell(r, 6, sym).font = Font(name=FONT, size=9, italic=True)
    c = put(ws, r, 7, formula, fmt, inp=inp)
    if res:
        c.font = F_R; c.fill = FILL_R
    ws.cell(r, 8, unit).font = F_N
    ws.cell(r, 9, note).font = F_S
    return f'$G${r}'


wb = openpyxl.Workbook()
# ======================= INPUT =======================
wi = wb.active; wi.title = 'Input'
for c, w in zip('ABCDEFGHI', (46, 4, 4, 4, 4, 10, 12, 9, 70)):
    wi.column_dimensions[c].width = w
wi['A1'] = 'ANALISIS DAYA DUKUNG & PENURUNAN TANAH DASAR - TIMBUNAN WELLPAD'; wi['A1'].font = F_T
wi['A2'] = 'Geotechnical Investigation & DED Geothermal Wellpad - PSPE Bora Pulu (PT EDC Panas Bumi Indonesia)'; wi['A2'].font = F_S
wi['A3'] = 'Sel kuning = input. Semua sheet BH dan Ringkasan dihitung dari sheet ini.'; wi['A3'].font = F_S
IN = {}
r = 5; sec(wi, r, 'A. BEBAN', 9)
IN['H'] = line(wi, 6, 'Tinggi timbunan', 'Hf', 2.0, 'm', 'arahan desain: timbunan 2 m di semua area', inp=True)
IN['gf'] = line(wi, 7, 'Berat isi material timbunan (dipadatkan)', 'γf', 18.0, 'kN/m³', 'asumsi timbunan pilihan dipadatkan', inp=True)
IN['qs'] = line(wi, 8, 'Beban merata di atas timbunan', 'qs', 100.0, 'kPa', 'arahan desain: 100 kPa di semua area', inp=True)
IN['q'] = line(wi, 9, 'Tekanan total pada tanah dasar', 'q', '=G6*G7+G8', 'kPa', 'q = Hf·γf + qs', res=True)
IN['B'] = line(wi, 10, 'Lebar area pembebanan', 'B', 40.0, 'm', 'ASUMSI ukuran platform/area beban - sesuaikan', inp=True)
IN['L'] = line(wi, 11, 'Panjang area pembebanan', 'L', 60.0, 'm', 'ASUMSI - sesuaikan (L ≥ B)', inp=True)
r = 13; sec(wi, r, 'B. KRITERIA', 9)
IN['SF'] = line(wi, 14, 'Faktor keamanan daya dukung', 'SF', 3.0, '-', 'SNI 8460:2017 (daya dukung fondasi dangkal)', inp=True)
IN['Hb'] = line(wi, 15, 'Kedalaman tinjauan kekuatan tanah daya dukung', 'Hb', 5.0, 'm', 'rata-rata parameter kuat geser 0 - Hb (sama dengan template)', inp=True)
IN['Slim'] = line(wi, 16, 'Batas penurunan total (informasi)', 'S izin', 100.0, 'mm', 'ASUMSI untuk platform timbunan - sesuaikan', inp=True)
r = 18; sec(wi, r, 'C. KORELASI PARAMETER (dari N-SPT bila tidak ada uji lab)', 9)
IN['kcu'] = line(wi, 19, 'Kuat geser undrained lempung/lanau: cu = k·N', 'k', 5.0, 'kPa/N', 'Stroud (1974), k ≈ 4,4-6; sama dengan PBCP & template', inp=True)
IN['ce'] = line(wi, 20, 'Koreksi energi N60 = CE·N', 'CE', 1.083, '-', 'sama dengan analisis likuefaksi (ER 65%)', inp=True)
IN['kEs'] = line(wi, 21, 'Modulus pasir: Es = k·N60', 'k', 1000.0, 'kPa', 'Kulhawy & Mayne (1990): E/pa = 10·N60 (pasir berlanau); sama dgn E50 = 1000N template', inp=True)
IN['kEu'] = line(wi, 22, 'Modulus undrained lempung/lanau: Eu = k·cu', 'k', 500.0, '-', 'Bowles (1996), Eu = 500-1500 cu (konservatif)', inp=True)
IN['rcr'] = line(wi, 23, 'Rasio Cr / Cc', '-', 0.2, '-', 'Cr ≈ 0,1-0,2 Cc (Das 2019)', inp=True)
IN['sup'] = line(wi, 24, "Rasio su / σ'p", '-', 0.22, '-', "Mesri (1975): σ'p = cu / 0,22", inp=True)
IN['Tv'] = line(wi, 25, 'Faktor waktu Tv untuk U = 90%', 'Tv', 0.848, '-', 'Terzaghi', inp=True)
IN['rinf'] = line(wi, 27, "Batas kedalaman pengaruh penurunan: Δσz/σ'v0 ≥", '-', 0.10, '-', "lapisan dengan Δσz < 10% σ'v0 diabaikan (Das 2019)", inp=True)
IN['nus'] = line(wi, 26, 'Poisson ratio pasir (drained)', 'ν', 0.3, '-', 'lanau/lempung undrained ν = 0,5', inp=True)
sec(wi, 29, 'D. MATERIAL PENGGANTI (REPLACEMENT)', 9)
IN['gr'] = line(wi, 30, 'Berat isi material pengganti (granular dipadatkan)', 'γr', 19.0, 'kN/m³', 'sirtu/selected fill dipadatkan', inp=True)
IN['phr'] = line(wi, 31, "Sudut geser material pengganti", "φ'r", 35.0, '°', 'granular dipadatkan', inp=True)
wi['A32'] = 'Kedalaman penggantian Dr diisi per borehole (sheet BH, sel G8).'; wi['A32'].font = F_S
wi['A34'] = "Cc = 0,009(LL-10) (Terzaghi & Peck 1967); φ' pasir = 27,1 + 0,3N60 - 0,00054N60² (Wolff 1989); φ' lanau/lempung: sin φ' = 0,8 - 0,094 ln PI (Terzaghi, Peck & Mesri 1996), dibatasi 20-30°."
wi['A35'] = 'Cv dari LL (NAVFAC DM-7.1, Fig. 5-7, contoh tak terganggu) - tabel di sheet Korelasi. Tidak ada uji konsolidasi (oedometer) pada penyelidikan ini.'
for a in ('A34', 'A35'):
    wi[a].font = F_S

# ======================= KORELASI =======================
wk = wb.create_sheet('Korelasi')
wk['A1'] = 'TABEL KORELASI Cv - LL (NAVFAC DM-7.1)'; wk['A1'].font = F_T
for j, h in enumerate(['LL (%)', 'Cv (cm²/s)', 'log10 Cv']):
    c = wk.cell(4, j + 1, h); c.font = F_H; c.fill = FILL_H; c.alignment = CEN; c.border = BOX
CV = [(30, 3e-3), (40, 1.5e-3), (50, 7e-4), (60, 4e-4), (80, 2e-4), (100, 1.2e-4), (120, 8e-5)]
for i, (ll, cv) in enumerate(CV):
    put(wk, 5 + i, 1, ll, '0', inp=True); put(wk, 5 + i, 2, cv, '0.0E+00', inp=True)
    put(wk, 5 + i, 3, f'=LOG10(B{5 + i})', '0.000')
for c, w in zip('ABC', (10, 12, 10)):
    wk.column_dimensions[c].width = w
LLT, LCV = 'Korelasi!$A$5:$A$11', 'Korelasi!$C$5:$C$11'

# ======================= SHEET PER BH =======================
lab = B.lab
SPEC = [  # (key, header, unit, width)
    ('no', 'No', '', 4), ('top', 'Atas', 'm', 6), ('bot', 'Bawah', 'm', 6), ('h', 'Tebal h', 'm', 6), ('z', 'Tengah z', 'm', 6),
    ('typ', 'Jenis', '', 6), ('N', 'N-SPT', '', 6), ('g', 'γ', 'kN/m³', 6), ('sv', 'σv bawah', 'kPa', 8), ('s0', "σ'v0 tengah", 'kPa', 8),
    ('LL', 'LL', '%', 6), ('PL', 'PL', '%', 6), ('PI', 'PI', '%', 6), ('e0', 'e0', '-', 6), ('cu', 'cu', 'kPa', 7), ('phi', "φ'", '°', 6),
    ('E', 'E', 'kPa', 9), ('nu', 'ν', '-', 5), ('m', 'm=B/2z', '-', 6), ('n', 'n=L/2z', '-', 6), ('I', 'I Bouss.', '-', 7),
    ('ds', 'Δσz pusat', 'kPa', 7), ('F1', 'F1 (bawah)', '-', 7), ('F2', 'F2 (bawah)', '-', 7), ('keep', 'faktor aktif', '-', 7),
    ('Si', 'Si', 'mm', 7), ('Cc', 'Cc', '-', 6), ('Cr', 'Cr', '-', 6), ('sp', "σ'p", 'kPa', 8), ('Sc', 'Sc', 'mm', 7),
    ('Cv', 'Cv', 'cm²/s', 9), ('t90', 't90', 'tahun', 7), ('hb0', 'Δ 0-Hb', 'm', 6), ('hbr', 'Δ 0-Hb sisa', 'm', 6), ('hzi', 'Δ 0-zI', 'm', 6),
    ('ket', 'Ket.', '', 34)]
COL = {k: CL(i + 1) for i, (k, *_ ) in enumerate(SPEC)}
summary = {}
for bh in BHS:
    ws = wb.create_sheet(bh)
    for i, (k, h, u, w) in enumerate(SPEC):
        ws.column_dimensions[CL(i + 1)].width = w
    NC = len(SPEC)
    g_un, g_sat = B.GAM[bh]
    ws['A1'] = f'ANALISIS DAYA DUKUNG, PENURUNAN SEGERA & KONSOLIDASI - {bh} ({B.AREA[bh]})'; ws['A1'].font = F_T
    ws['A2'] = 'PSPE Bora Pulu - timbunan + beban merata (lihat sheet Input)'; ws['A2'].font = F_S
    sec(ws, 4, '1. PARAMETER BOREHOLE', NC)
    DR0 = {'BH-01': 3.0, 'BH-04': 3.0}.get(bh, 0.0)
    meta = [('Muka air tanah (GWL)', B.GWL[bh], 'm', 'GWL analisis likuefaksi (Report BP-00-GEO-RPT-003)'),
            ('Berat isi di atas GWL', g_un, 'kN/m³', 'Report Tabel 5.1 (hasil lab)'),
            ('Berat isi di bawah GWL', g_sat, 'kN/m³', 'Report Tabel 5.1 (hasil lab)'),
            ('Kedalaman penggantian tanah (replacement) - kondisi desain', DR0, 'm',
             '0 = tanpa perbaikan; tanah 0 - Dr diganti material granular (sheet Input bagian D)')]
    for k, (a, v, u, n) in enumerate(meta):
        ws.cell(5 + k, 1, a).font = F_N
        put(ws, 5 + k, 7, v, '0.0', inp=True); ws.cell(5 + k, 8, u).font = F_N; ws.cell(5 + k, 9, n).font = F_S
    GW, G1, G2, DR = '$G$5', '$G$6', '$G$7', '$G$8'
    q, Bw, Lw = 'Input!' + IN['q'], 'Input!' + IN['B'], 'Input!' + IN['L']
    HB = 'Input!' + IN['Hb']

    sec(ws, 10, '2. TABEL TANAH & PERHITUNGAN PER LAPISAN (sublapisan = interval SPT 1,5 m)', NC)
    HR = 11
    for j, (k, h, u, w) in enumerate(SPEC):
        for rr, t in ((HR, h), (HR + 1, u)):
            c = ws.cell(rr, j + 1, t); c.font = F_H; c.fill = FILL_H; c.alignment = CEN; c.border = BOX
    ws.row_dimensions[HR].height = 26
    sub = B.merge(B.sublayers(bh), 41)
    fd = B.fine_default(bh)
    R0 = HR + 2; RN = R0 + len(sub) - 1
    ZIROW = RN + 60          # diisi ulang di bawah
    X = COL
    for i, s in enumerate(sub):
        r = R0 + i
        c = lambda k: f'{X[k]}{r}'
        p = lambda k: f'{X[k]}{r - 1}'
        ls = B.labstats(bh, s['top'], s['bot'])
        fine = s['cat'] != 'Sand'
        note = []
        if s['src'] != bh:
            note.append(f'data salinan {s["src"]}')
        if i == 0:
            note.append('0-1,5 m tidak di-SPT: N lapisan berikut')
        if fine and ls['LL'] is None:
            note.append('LL/PL/e0 = rata-rata BH')
        LL = (ls['LL'] or fd['LL']) if fine else None
        PL = (ls['PL'] or fd['PL']) if fine else None
        e0 = ls['e'] or (fd['e'] if fine else 0.65)
        sand = f'{c("typ")}="Sand"'
        def F12(zref):
            M = f'({Lw}/{Bw})'; Nn = f'({zref}/({Bw}/2))'
            s1 = f'SQRT({M}^2+{Nn}^2)'; s2 = f'SQRT({M}^2+{Nn}^2+1)'; s3 = f'SQRT({M}^2+1)'
            f1 = (f'({M}*LN((1+{s3})*{s1}/({M}*(1+{s2})))+LN(({M}+{s3})*SQRT(1+{Nn}^2)/({M}+{s2})))/PI()')
            f2 = f'{Nn}/(2*PI())*ATAN({M}/({Nn}*{s2}))'
            return f1, f2
        f1, f2 = F12(c('bot'))
        cnu = f'(1-2*{c("nu")})/(1-{c("nu")})'
        F1t = '0' if i == 0 else p('F1'); F2t = '0' if i == 0 else p('F2')
        vals = {
            'no': i + 1, 'top': s['top'], 'bot': s['bot'], 'h': f'={c("bot")}-{c("top")}', 'z': f'=({c("top")}+{c("bot")})/2',
            'typ': s['cat'], 'N': s['N'],
            'g': f'=IF({c("z")}>{GW},{G2},{G1})',
            'sv': f'={c("g")}*{c("h")}' if i == 0 else f'={p("sv")}+{c("g")}*{c("h")}',
            's0': f'={c("sv")}-{c("g")}*{c("h")}/2-10*MAX({c("z")}-{GW},0)',
            'LL': LL, 'PL': PL, 'PI': f'=IF({c("LL")}="","",{c("LL")}-{c("PL")})', 'e0': round(e0, 3),
            'cu': f'=IF({sand},0,Input!{IN["kcu"]}*{c("N")})',
            'phi': (f'=IF({sand},27.1+0.3*MIN(Input!{IN["ce"]}*{c("N")},60)-0.00054*MIN(Input!{IN["ce"]}*{c("N")},60)^2,'
                    f'MIN(MAX(DEGREES(ASIN(0.8-0.094*LN(MAX(IF({c("PI")}="",10,{c("PI")}),5)))),20),30))'),
            'E': f'=IF({sand},Input!{IN["kEs"]}*Input!{IN["ce"]}*{c("N")},Input!{IN["kEu"]}*{c("cu")})',
            'nu': f'=IF({sand},Input!{IN["nus"]},0.5)',
            'm': f'={Bw}/2/{c("z")}', 'n': f'={Lw}/2/{c("z")}',
            'I': (f'=1/(4*PI())*(2*{c("m")}*{c("n")}*SQRT({c("m")}^2+{c("n")}^2+1)/({c("m")}^2+{c("n")}^2+1+{c("m")}^2*{c("n")}^2)'
                  f'*({c("m")}^2+{c("n")}^2+2)/({c("m")}^2+{c("n")}^2+1)'
                  f'+ATAN2({c("m")}^2+{c("n")}^2+1-{c("m")}^2*{c("n")}^2,2*{c("m")}*{c("n")}*SQRT({c("m")}^2+{c("n")}^2+1)))'),
            'ds': f'=4*{q}*{c("I")}',
            'F1': '=' + f1, 'F2': '=' + f2,
            'keep': f'=MAX(0,{c("bot")}-MAX({c("top")},{DR}))/{c("h")}*IF({c("ds")}>=Input!{IN["rinf"]}*{c("s0")},1,0)',
            'Si': (f'=4*{q}*({Bw}/2)*(1-{c("nu")}^2)/{c("E")}*(({c("F1")}-{F1t})+{cnu}*({c("F2")}-{F2t}))*1000*{c("keep")}'),
            'Cc': f'=IF({sand},0,0.009*({c("LL")}-10))',
            'Cr': f'={c("Cc")}*Input!{IN["rcr"]}',
            'sp': f'=IF({sand},0,MAX({c("cu")}/Input!{IN["sup"]},{c("s0")}))',
            'Sc': (f'=IF({sand},0,1000*{c("h")}/(1+{c("e0")})*IF({c("s0")}+{c("ds")}<={c("sp")},{c("Cr")}*LOG10(({c("s0")}+{c("ds")})/{c("s0")}),'
                   f'{c("Cr")}*LOG10({c("sp")}/{c("s0")})+{c("Cc")}*LOG10(({c("s0")}+{c("ds")})/{c("sp")})))*{c("keep")}'),
            'Cv': (f'=IF({sand},0,10^(INDEX({LCV},MIN(MATCH(MIN(MAX({c("LL")},30),120),{LLT},1),6))'
                   f'+(MIN(MAX({c("LL")},30),120)-INDEX({LLT},MIN(MATCH(MIN(MAX({c("LL")},30),120),{LLT},1),6)))'
                   f'/(INDEX({LLT},MIN(MATCH(MIN(MAX({c("LL")},30),120),{LLT},1),6)+1)-INDEX({LLT},MIN(MATCH(MIN(MAX({c("LL")},30),120),{LLT},1),6)))'
                   f'*(INDEX({LCV},MIN(MATCH(MIN(MAX({c("LL")},30),120),{LLT},1),6)+1)-INDEX({LCV},MIN(MATCH(MIN(MAX({c("LL")},30),120),{LLT},1),6)))))'),
            't90': f'=IF(OR({sand},{c("keep")}=0),0,Input!{IN["Tv"]}*({c("h")}/2)^2/({c("Cv")}*1E-4*31536000))',
            'hb0': f'=MAX(0,MIN({c("bot")},{HB})-{c("top")})',
            'hbr': f'=MAX(0,MIN({c("bot")},{HB})-MAX({c("top")},{DR}))',
            'hzi': f'=MAX(0,MIN({c("bot")},$G$ZI)-{c("top")})',
            'ket': '; '.join(note),
        }
        fmts = {'no': '0', 'top': '0.00', 'bot': '0.00', 'h': '0.00', 'z': '0.00', 'N': '0', 'g': '0.0', 'sv': '0.0', 's0': '0.0',
                'LL': '0.0', 'PL': '0.0', 'PI': '0.0', 'e0': '0.000', 'cu': '0;-0;"-"', 'phi': '0.0', 'E': '#,##0', 'nu': '0.00',
                'm': '0.00', 'n': '0.00', 'I': '0.000', 'ds': '0.0', 'F1': '0.0000', 'F2': '0.0000', 'keep': '0.00', 'Si': '0.0',
                'Cc': '0.000;-0;"-"', 'Cr': '0.000;-0;"-"', 'sp': '0;-0;"-"', 'Sc': '0.0;-0;"-"', 'Cv': '0.0E+00;-0;"-"',
                't90': '0.00;-0;"-"', 'hb0': '0.00', 'hbr': '0.00', 'hzi': '0.00'}
        for j, (k, *_ ) in enumerate(SPEC):
            cell = put(ws, r, j + 1, vals[k], fmts.get(k), inp=k in ('top', 'bot', 'typ', 'N', 'LL', 'PL', 'e0'))
            cell.alignment = Alignment(horizontal='left' if k == 'ket' else 'center')
    rng = lambda k: f'{X[k]}{R0}:{X[k]}{RN}'
    ws.cell(RN + 1, 1, 'Sel kuning = data (SPT, lab, laporan). Jenis: Sand / Silt / Clay. Sublapisan di bawah kedalaman bor memakai data borehole terdekat (lihat Ket.).').font = F_S

    # -------- 3. DAYA DUKUNG --------
    r = RN + 3
    sec(ws, r, '3. DAYA DUKUNG TANAH DASAR (Terzaghi/Vesic, SNI 8460:2017)', NC); r += 1
    ws.cell(r, 1, 'Kondisi undrained (jangka pendek) dihitung bila ada lanau/lempung dalam 0 - Hb; kondisi drained (jangka panjang) selalu dihitung; '
                  'daya dukung = nilai terkecil. Tanah 0 - Dr yang diganti dianggap material granular (Input bagian D).').font = F_S; r += 1
    res = {}
    for tag, title, dr in (('0', 'a) TANPA PERBAIKAN', '0'), ('r', 'b) DENGAN PENGGANTIAN TANAH 0 - Dr (kondisi desain)', DR)):
        ws.cell(r, 1, title).font = F_B; r += 1
        ovl = rng('hb0') if tag == '0' else rng('hbr')
        drh = f'MIN({dr},{HB})'
        hf = line(ws, r, 'Tebal lanau/lempung asli dalam 0 - Hb', 'Σh', f'=SUMPRODUCT(--({rng("typ")}<>"Sand"),{ovl})', 'm'); r += 1
        cu = line(ws, r, 'cu rata-rata lanau/lempung (0 - Hb)', 'cu', f'=IF({hf}=0,0,SUMPRODUCT(--({rng("typ")}<>"Sand"),{ovl},{rng("cu")})/{hf})', 'kPa', 'cu = k·N'); r += 1
        phi = line(ws, r, "φ' rata-rata (0 - Hb)", "φ'", f'=(SUMPRODUCT({ovl},{rng("phi")})+{drh}*Input!{IN["phr"]})/{HB}', '°',
                   "lapisan diganti memakai φ' material pengganti"); r += 1
        gam = line(ws, r, "Berat isi efektif rata-rata (0 - Hb)", "γ'",
                   f'=(SUMPRODUCT({ovl},{rng("g")})+{drh}*Input!{IN["gr"]})/{HB}-10*MAX(0,{HB}-{GW})/{HB}', 'kN/m³'); r += 1
        sc = line(ws, r, 'Undrained: faktor bentuk (Nc = 5,14)', 'sc', f'=1+0.2*{Bw}/{Lw}', '-'); r += 1
        quu = line(ws, r, 'qu undrained = cu·Nc·sc', 'qu,u', f'=IF({hf}=0,"-",{cu}*5.14*{sc})', 'kPa', 'tidak berlaku bila tidak ada lanau/lempung', '0.0'); r += 1
        nq = line(ws, r, 'Drained: Nq = e^(π tanφ)·tan²(45+φ/2)', 'Nq', f'=EXP(PI()*TAN(RADIANS({phi})))*TAN(RADIANS(45+{phi}/2))^2', '-'); r += 1
        ng = line(ws, r, 'Nγ = 2(Nq+1)tanφ (Vesic 1975)', 'Nγ', f'=2*({nq}+1)*TAN(RADIANS({phi}))', '-'); r += 1
        sg = line(ws, r, 'sγ = 1 - 0,4B/L ;  rγ = 1 - 0,25log(B/2) (Bowles 1996)', 'sγ·rγ', f'=(1-0.4*{Bw}/{Lw})*IF({Bw}>2,1-0.25*LOG10({Bw}/2),1)', '-'); r += 1
        qud = line(ws, r, "qu drained = 0,5·γ'·B·Nγ·sγ·rγ   (c' = 0, q = 0)", 'qu,d', f'=0.5*{gam}*{Bw}*{ng}*{sg}', 'kPa', '', '0.0'); r += 1
        qu = line(ws, r, 'Daya dukung ultimit menentukan', 'qu', f'=IF({hf}=0,{qud},MIN({quu},{qud}))', 'kPa', '', '0.0', res=True); r += 1
        qa = line(ws, r, 'Daya dukung izin = qu / SF', 'qall', f'={qu}/Input!{IN["SF"]}', 'kPa', '', '0.0', res=True); r += 1
        line(ws, r, 'Tekanan kerja (timbunan + beban merata)', 'q', f'={q}', 'kPa', '', '0.0'); r += 1
        chk = line(ws, r, 'Kontrol daya dukung', '', f'=IF({qa}>={q},"qall ≥ q … OK","qall < q … NG")', '', '', '@', res=True); r += 1
        sfa = line(ws, r, 'Faktor keamanan aktual = qu / q', 'FS', f'={qu}/{q}', '-'); r += 1
        hall = line(ws, r, 'Tinggi timbunan izin (dengan beban merata)', 'H izin', f'=MAX(0,({qa}-Input!{IN["qs"]})/Input!{IN["gf"]})', 'm', 'H izin = (qall - qs)/γf'); r += 2
        res[tag] = dict(hf=hf, cu=cu, phi=phi, quu=quu, qud=qud, qa=qa, chk=chk, fs=sfa, hall=hall)
    ws.cell(r, 1, 'c) Tanah pasiran - daya dukung izin berbasis penurunan 25 mm (Meyerhof 1965, Bowles 1996)').font = F_B; r += 1
    nb = line(ws, r, 'N60 rata-rata lapisan pasir (0 - Hb)', 'N60',
              f'=IFERROR(SUMPRODUCT(--({rng("typ")}="Sand"),{rng("hb0")},{rng("N")})/SUMPRODUCT(--({rng("typ")}="Sand"),{rng("hb0")})*Input!{IN["ce"]},"-")', '-', '', '0.0'); r += 1
    qam = line(ws, r, 'qa (S = 25 mm) = (N55/0,08)·((B+0,3)/B)²·Kd', 'qa', f'=IFERROR(({nb}*60/55)/0.08*(({Bw}+0.3)/{Bw})^2,"-")', 'kPa', 'Kd = 1 (beban di permukaan)', '0.0'); r += 1
    ws.cell(r, 1, 'Pada tanah pasiran daya dukung geser hampir selalu jauh di atas beban; penentu desain adalah penurunan (bagian 4).').font = F_S; r += 2

    # -------- 4. PENURUNAN SEGERA --------
    sec(ws, r, '4. PENURUNAN SEGERA (IMMEDIATE SETTLEMENT) - kondisi desain', NC); r += 1
    ws.cell(r, 1, "d) Teori elastis berlapis (Steinbrenner 1934; Bowles 1996): pusat area fleksibel B×L, tiap lapisan Si = 4q(B/2)(1-ν²)/E·[ΔF1 + (1-2ν)/(1-ν)·ΔF2]."
                  ' Pasir: E drained, ν pasir; lanau/lempung: Eu undrained, ν = 0,5.').font = F_S; r += 1
    si = line(ws, r, 'Penurunan segera total (dalam kedalaman pengaruh)', 'Si', f'=SUM({rng("Si")})', 'mm', '', '0.0', res=True); r += 1
    si_s = line(ws, r, '   - dari lapisan pasir', '', f'=SUMPRODUCT(--({rng("typ")}="Sand"),{rng("Si")})', 'mm', '', '0.0'); r += 1
    line(ws, r, '   - dari lapisan lanau/lempung (undrained)', '', f'={si}-{si_s}', 'mm', '', '0.0'); r += 2
    ws.cell(r, 1, 'e) Pemeriksaan tanah pasiran - Burland & Burbidge (1985), pasir NC:  S = fs·fl·q·B^0,7·Ic,   Ic = 1,71/N̄^1,4').font = F_B; r += 1
    zi = line(ws, r, 'Kedalaman pengaruh zI = B^0,763', 'zI', f'=MIN({Bw}^0.763,{X["bot"]}{RN})', 'm'); ZIROW = r; r += 1
    nbar = line(ws, r, 'N rata-rata dalam zI (N terukur)', 'N̄', f'=SUMPRODUCT({rng("hzi")},{rng("N")})/{zi}', '-', '', '0.0'); r += 1
    ic = line(ws, r, 'Indeks kompresibilitas', 'Ic', f'=1.71/{nbar}^1.4', '-', '', '0.0000'); r += 1
    fs_ = line(ws, r, 'Faktor bentuk fs = [1,25(L/B)/(L/B+0,25)]²', 'fs', f'=(1.25*{Lw}/{Bw}/({Lw}/{Bw}+0.25))^2', '-'); r += 1
    sbb = line(ws, r, 'Penurunan (B&B), fl = 1', 'S', f'={fs_}*{q}*{Bw}^0.7*{ic}', 'mm', '', '0.0', res=True); r += 1
    ws.cell(r, 1, 'B&B menganggap seluruh zona pengaruh berupa pasir; bila ada lanau/lempung, pakai butir d) sebagai nilai desain.').font = F_S; r += 2

    # -------- 5. KONSOLIDASI --------
    sec(ws, r, '5. PENURUNAN KONSOLIDASI PRIMER & WAKTU KONSOLIDASI (lapisan lanau/lempung) - kondisi desain', NC); r += 1
    ws.cell(r, 1, "Sc = h/(1+e0)·[Cr·log(σ'p/σ'0) + Cc·log((σ'0+Δσ)/σ'p)] (Das 2019); Δσ Boussinesq di pusat area; σ'p = max(cu/0,22 ; σ'0); "
                  't90 = Tv·Hdr²/Cv, drainase ganda (Hdr = h/2) karena diapit pasir.').font = F_S; r += 1
    scs = line(ws, r, 'Penurunan konsolidasi total', 'Sc', f'=SUM({rng("Sc")})', 'mm', '', '0.0', res=True); r += 1
    t90 = line(ws, r, 'Waktu konsolidasi 90% (lapisan terlama)', 't90', f'=MAX({rng("t90")})', 'tahun', '', '0.00'); r += 2
    sec(ws, r, '6. RINGKASAN', NC); r += 1
    stot = line(ws, r, 'Penurunan total jangka panjang = Si + Sc', 'S', f'={si}+{scs}', 'mm', '', '0.0', res=True); r += 1
    line(ws, r, 'Kontrol terhadap batas penurunan', '', f'=IF({stot}<=Input!{IN["Slim"]},"S ≤ S izin … OK","S > S izin … NG")', '', '', '@', res=True); r += 1
    ws.cell(r, 1, 'Penurunan pasca-likuefaksi saat gempa (laporan likuefaksi, Tabel 5.4: 100 - 383 mm) TIDAK termasuk; ditinjau terpisah.').font = F_S
    for rr in range(R0, RN + 1):
        cc = ws[f'{X["hzi"]}{rr}']; cc.value = cc.value.replace('$G$ZI', f'$G${ZIROW}')
    ch = ScatterChart(); ch.title = f'Penurunan per sublapisan - {bh}'; ch.style = 13
    ch.y_axis.title = 'Kedalaman tengah (m)'; ch.x_axis.title = 'Penurunan (mm)'; ch.y_axis.scaling.orientation = 'maxMin'
    ch.x_axis.delete = False; ch.y_axis.delete = False
    for k, nm in (('Si', 'Si segera'), ('Sc', 'Sc konsolidasi')):
        ci = list(COL).index(k) + 1
        sr = Series(Reference(ws, min_col=5, min_row=R0, max_row=RN), Reference(ws, min_col=ci, min_row=R0, max_row=RN), title=nm)
        sr.marker.symbol = 'circle'; sr.marker.size = 4
        ch.series.append(sr)
    ch.height = 10; ch.width = 15
    ws.add_chart(ch, f'M{RN + 4}')
    ws.freeze_panes = f'B{R0}'
    ws.sheet_view.zoomScale = 85
    ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    summary[bh] = dict(res=res, qam=qam, si=si, sbb=sbb, sc=scs, t90=t90, st=stot, dr=DR)

# ======================= RINGKASAN =======================
wr = wb.create_sheet('Ringkasan', 1)
wr['A1'] = 'RINGKASAN DAYA DUKUNG & PENURUNAN TANAH DASAR'; wr['A1'].font = F_T
wr['A2'] = '=" Beban: timbunan "&TEXT(Input!G6,"0.0")&" m + beban merata "&TEXT(Input!G8,"0")&" kPa  →  q = "&TEXT(Input!G9,"0")&" kPa;  area "&TEXT(Input!G10,"0")&" × "&TEXT(Input!G11,"0")&" m"'
wr['A2'].font = F_S
hd = ['Borehole', 'Area', 'Lanau/lempung 0-Hb (m)', 'cu (kPa)', "φ' (°)", 'qall tanpa perbaikan (kPa)', 'FS tanpa perbaikan', 'Kontrol',
      'Dr replacement (m)', 'qall desain (kPa)', 'FS desain', 'Kontrol desain', 'H timbunan izin desain (m)', 'qa pasir S=25 mm (kPa)',
      'Si elastis (mm)', 'Si pasir B&B (mm)', 'Sc (mm)', 't90 (tahun)', 'S total (mm)']
for j, h in enumerate(hd):
    c = wr.cell(4, j + 1, h); c.font = F_H; c.fill = FILL_H; c.alignment = CEN; c.border = BOX
    wr.column_dimensions[CL(j + 1)].width = 11 if j > 1 else 10
wr.row_dimensions[4].height = 44
for i, bh in enumerate(BHS):
    r = 5 + i; S_ = summary[bh]; r0, rr_ = S_['res']['0'], S_['res']['r']
    refs = [r0['hf'], r0['cu'], r0['phi'], r0['qa'], r0['fs'], r0['chk'], S_['dr'], rr_['qa'], rr_['fs'], rr_['chk'], rr_['hall'],
            S_['qam'], S_['si'], S_['sbb'], S_['sc'], S_['t90'], S_['st']]
    put(wr, r, 1, bh, bold=True); put(wr, r, 2, B.AREA[bh])
    for j, ref in enumerate(refs):
        fmt = '@' if ref in (r0['chk'], rr_['chk']) else ('0.00' if ref in (r0['fs'], rr_['fs'], S_['t90']) else '0.0')
        put(wr, r, 3 + j, f"='{bh}'!{ref}", fmt).alignment = Alignment(horizontal='center')
notes = [
    'CATATAN',
    '1. Daya dukung: undrained (cu·Nc) bila ada lanau/lempung dalam 0 - Hb; drained (Vesic) untuk pasir. Pada area beban lebar (B besar) suku Nγ membuat qu drained sangat besar -',
    '   artinya keruntuhan geser tidak menentukan; yang menentukan adalah penurunan.',
    '2. Cara analisis tanah pasiran: (a) daya dukung drained dengan φ\' dari N-SPT; (b) qa berbasis penurunan 25 mm (Meyerhof/Bowles); (c) penurunan segera dengan',
    '   teori elastis berlapis (E dari N-SPT, batas atas konservatif karena E dianggap konstan) dan Burland & Burbidge (empiris, lebih realistis untuk pasir);',
    '   (d) pasir tidak mengalami konsolidasi primer; (e) pasir jenuh lepas: tambahkan penurunan pasca-likuefaksi pada kondisi gempa.',
    '3. Ukuran area beban B×L di sheet Input adalah ASUMSI (40 × 60 m). Penurunan sangat dipengaruhi ukuran area: untuk B = L = 10 m, Si turun ±50-75%.',
    '4. Penurunan pasca-likuefaksi saat gempa (laporan likuefaksi Tabel 5.4): BH-01 289, BH-02 383, BH-03 227, BH-04 159, BH-05 100 mm - ditinjau terpisah.',
    '5. BH-01 kedalaman 1,5 m: data SPT = lempung N5, laporan likuefaksi/lab = pasir SP N12. Dipakai data SPT (konservatif).',
]
for k, t in enumerate(notes):
    wr.cell(11 + k, 1, t).font = F_B if k == 0 else F_S
for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False
wb.calculation.fullCalcOnLoad = True
OUT = 'out/Analisis_DayaDukung_Penurunan_BoraPulu.xlsx'
wb.save(OUT)
print(OUT)
