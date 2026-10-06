"""Bearing capacity, immediate & consolidation settlement - PSPE Bora Pulu.
A4 calculation-sheet layout following the client's template (Proxima Nova Lt 10.5, narrow grid B:AF, framed pages)."""
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.worksheet.pagebreak import Break, RowBreak
from openpyxl.drawing.image import Image as XLImage
import math
import build_bc as B

FN, FS = 'Proxima Nova Lt', 10.5
f_n = Font(name=FN, size=FS)
f_b = Font(name=FN, size=FS, bold=True)
f_i = Font(name=FN, size=FS, italic=True)
f_ok = Font(name=FN, size=FS, bold=True, color='002060')
f_note = Font(name=FN, size=9.5, italic=True)
FILL_IN = PatternFill('solid', fgColor='CCFFFF')      # template input colour
FILL_HD = PatternFill('solid', fgColor='FFFFCC')      # template table-header colour
th, md, hr = Side(style='thin'), Side(style='medium'), Side(style='hair')
C_L, C_C, C_R = Alignment(horizontal='left', vertical='center'), Alignment(horizontal='center', vertical='center', wrap_text=True), Alignment(horizontal='right', vertical='center')
PAGE = 72                     # rows per printed A4 page (scale 66 %)
FIRST, LASTC = 3, 31          # content columns C..AE ; B and AF are frame margins
BHS = ['BH-01', 'BH-02', 'BH-03', 'BH-04', 'BH-05']
DATE = datetime.datetime(2026, 10, 6)
PROJECT = 'GEOTHERMAL WELLPAD PSPE BORA PULU - SIGI, CENTRAL SULAWESI'
CLIENT = 'PT EDC PANAS BUMI INDONESIA'


def setup(ws):
    widths = {1: 3.43, 2: 4.0, 10: 3.86, 11: 4.0, 32: 3.43}
    for c in range(1, 33):
        ws.column_dimensions[CL(c)].width = widths.get(c, 3.57 if 3 <= c <= 9 else 4.57)
    ws.sheet_format.defaultRowHeight = 13.9
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ps = ws.page_setup
    ps.paperSize = 9; ps.orientation = 'portrait'; ps.scale = 66; ps.fitToHeight = 0
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.page_margins.top = ws.page_margins.bottom = 0.6
    ws.print_options.horizontalCentered = True


def W(ws, r, c1, c2, v, font=f_n, fmt=None, fill=None, align=C_L, border=None):
    if c2 > c1:
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    cell = ws.cell(r, c1, v)
    cell.font = font; cell.alignment = align
    if fmt:
        cell.number_format = fmt
    for c in range(c1, c2 + 1):
        if fill:
            ws.cell(r, c).fill = fill
        if border:
            ws.cell(r, c).border = border
    return cell


def text(ws, r, s, font=f_n, c1=FIRST, c2=LASTC):
    return W(ws, r, c1, c2, s, font)


class Page:
    """framed A4 page with template header block and sheet number"""
    def __init__(self, ws, title, location):
        self.ws, self.title, self.loc, self.n = ws, title, location, 0
        self.p0 = None

    def new(self, used=None):
        ws = self.ws
        if self.p0 is not None and used is not None:
            assert used <= self.end(), (self.title, self.n, used, self.end())
        self.n += 1
        self.p0 = 2 + (self.n - 1) * PAGE
        p0 = self.p0
        box = Border(left=th, right=th, top=th, bottom=th)
        rows = [('PROJECT :', PROJECT, 'Prepared by :', 'Date'), ('CLIENT :', CLIENT, 'DF', DATE),
                ('LOCATION :', self.loc, 'Verified by :', 'Date'), ('SHEET :', self.title, 'BC', DATE)]
        for k, (a, b, c, d) in enumerate(rows):
            r = p0 + k
            W(ws, r, 3, 7, a, f_b, align=C_R)
            W(ws, r, 8, 23, b, f_b)
            W(ws, r, 24, 27, c, f_b, align=C_C, border=box)
            W(ws, r, 28, 31, d, f_b, fmt='dd-mm-yyyy', align=C_C, border=box)
        for c in range(2, 33):
            cell = ws.cell(p0 + 3, c)
            cell.border = Border(left=cell.border.left, right=cell.border.right, top=cell.border.top, bottom=Side(style='double'))
        p1 = p0 + PAGE - 1
        W(ws, p1, 24, 27, 'Sheet :', f_b, align=C_C)
        W(ws, p1, 28, 31, self.n, f_b, align=C_C)
        # frame
        for r in range(p0, p1 + 1):
            for c, side in ((2, 'left'), (32, 'right')):
                cell = ws.cell(r, c); b = cell.border
                cell.border = Border(left=md if side == 'left' else b.left, right=md if side == 'right' else b.right, top=b.top, bottom=b.bottom)
        for c in range(2, 33):
            for r, side in ((p0, 'top'), (p1, 'bottom')):
                cell = ws.cell(r, c); b = cell.border
                cell.border = Border(left=b.left, right=b.right, top=md if side == 'top' else b.top, bottom=md if side == 'bottom' else b.bottom)
        if self.n > 1:
            ws.row_breaks.append(Break(id=p0 - 1))
        return p0 + 5            # first free row

    def end(self):
        return self.p0 + PAGE - 2    # last usable row

    def close(self, used=None):
        if used is not None:
            assert used <= self.end(), (self.title, self.n, used, self.end())
        last = 2 + self.n * PAGE - 1
        self.ws.print_area = f'B2:AF{last}'


def figure(ws, r, name, caption, width=860):
    img = XLImage(f'bc/fig/{name}.png')
    w0, h0 = img.width, img.height
    img.width = width; img.height = int(h0 * width / w0)
    ws.add_image(img, f'C{r}')
    r += math.ceil(img.height / 18.5) + 1
    W(ws, r, 3, 31, caption, f_b, align=Alignment(horizontal='center', vertical='center'))
    return r + 2


def calc(ws, r, label, eq, sym, val, unit='', fmt='0.00', kind='n'):
    """template-style calculation line: label | equation | symbol = | value | unit"""
    W(ws, r, 3, 14, label)
    W(ws, r, 15, 23, eq, f_i)
    W(ws, r, 24, 26, sym, align=C_C, border=Border(left=th, right=th, top=th, bottom=th))
    fill = FILL_IN if kind == 'in' else None
    font = f_ok if kind == 'res' else f_n
    c = W(ws, r, 27, 29, val, font, fmt, fill, C_C, Border(left=th, right=th, top=th, bottom=th))
    W(ws, r, 30, 31, unit)
    return f'$AA${r}'


def table_header(ws, r, cols, h2=None):
    """cols: list of (title, unit, width_in_grid_cols); returns list of (c1, c2)"""
    spans, c = [], FIRST
    for t, u, w in cols:
        spans.append((c, c + w - 1))
        bd = Border(left=th, right=th, top=md, bottom=th)
        W(ws, r, c, c + w - 1, t, f_b, fill=FILL_HD, align=C_C, border=bd)
        W(ws, r + 1, c, c + w - 1, u, f_b, fill=FILL_HD, align=C_C, border=Border(left=th, right=th, top=th, bottom=th))
        c += w
    ws.row_dimensions[r].height = 27
    return spans


def table_row(ws, r, spans, vals, fmts, inputs=(), last=False):
    for k, ((c1, c2), v, fm) in enumerate(zip(spans, vals, fmts)):
        bd = Border(left=th, right=th, top=hr, bottom=md if last else hr)
        W(ws, r, c1, c2, v, f_n, fm, FILL_IN if k in inputs else None, C_C, bd)


# =====================================================================================
wb = openpyxl.Workbook()
wi = wb.active; wi.title = 'Input'
setup(wi)
P = Page(wi, 'DESIGN INPUT & CORRELATIONS', 'All boreholes (BH-01 to BH-05)')
r = P.new(r if 'r' in dir() else None)
IN = {}
text(wi, r, '1. DESIGN INPUT', f_b); r += 1
text(wi, r, 'Input cells are shaded blue. All borehole sheets and the summary are calculated from this sheet.', f_note); r += 2
text(wi, r, '1.1. Design load', f_b); r += 1
IN['H'] = calc(wi, r, 'Height of embankment fill', 'design brief', 'Hf =', 2.0, 'm', kind='in'); r += 1
IN['gf'] = calc(wi, r, 'Unit weight of compacted fill', 'assumed', 'γf =', 18.0, 'kN/m³', kind='in'); r += 1
IN['qs'] = calc(wi, r, 'Uniform surcharge on platform', 'design brief', 'qs =', 100.0, 'kPa', kind='in'); r += 1
IN['q'] = calc(wi, r, 'Total pressure on subgrade', 'q = Hf·γf + qs', 'q =', f"={IN['H']}*{IN['gf']}+{IN['qs']}", 'kPa', '0.0', 'res'); r += 1
IN['B'] = calc(wi, r, 'Width of embankment area', 'design brief', 'B =', 74.0, 'm', kind='in'); r += 1
IN['L'] = calc(wi, r, 'Length of embankment area (L ≥ B)', 'design brief', 'L =', 202.0, 'm', kind='in'); r += 2
text(wi, r, '1.2. Design criteria', f_b); r += 1
IN['SF'] = calc(wi, r, 'Factor of safety for bearing capacity', 'SNI 8460:2017', 'SF =', 3.0, '-', kind='in'); r += 1
IN['Hb'] = calc(wi, r, 'Depth for averaging shear strength', 'as previous template', 'Hb =', 5.0, 'm', kind='in'); r += 1
IN['rinf'] = calc(wi, r, "Influence depth limit: Δσz / σ'v0 ≥", 'dense / OC soil', 'ratio =', 0.20, '-', kind='in'); r += 1
IN['Slim'] = calc(wi, r, 'Allowable total settlement (assumed)', 'to be confirmed', 'S all =', 100.0, 'mm', kind='in'); r += 1
IN['slope'] = calc(wi, r, 'Embankment side slope, horizontal : 1 vertical', '1V : nH', 'n =', 2.0, '-', '0.0', 'in'); r += 1
IN['FSsq'] = calc(wi, r, 'Required FS against lateral squeeze', 'FHWA-NHI-06-088', 'FS req =', 1.5, '-', kind='in'); r += 1
IN['ncmax'] = calc(wi, r, 'Upper limit of squeezing factor Nc* (0 = no limit)', 'optional, conservative', 'Nc*max =', 0.0, '-', '0.0', 'in'); r += 2
text(wi, r, '1.3. Replacement material', f_b); r += 1
IN['gr'] = calc(wi, r, 'Unit weight of compacted granular fill', 'assumed', 'γr =', 19.0, 'kN/m³', kind='in'); r += 1
IN['phr'] = calc(wi, r, 'Friction angle of granular fill', 'assumed', "φ'r =", 35.0, '°', kind='in'); r += 1
text(wi, r, 'Replacement depth Dr is entered on each borehole sheet.', f_note); r += 2
text(wi, r, '2. CORRELATIONS (used where no laboratory test is available)', f_b); r += 2
IN['kcu'] = calc(wi, r, 'Undrained shear strength, cu = k·N', 'Stroud (1974)', 'k =', 5.0, 'kPa', kind='in'); r += 1
IN['ce'] = calc(wi, r, 'Energy correction, N60 = CE·N', 'liquefaction report', 'CE =', 1.083, '-', '0.000', 'in'); r += 1
IN['kEs'] = calc(wi, r, 'Sand modulus, Es = k·N60', 'Kulhawy & Mayne (1990)', 'k =', 1000.0, 'kPa', '0', 'in'); r += 1
IN['kEu'] = calc(wi, r, 'Silt/clay undrained modulus, Eu = k·cu', 'Bowles (1996)', 'k =', 500.0, '-', '0', 'in'); r += 1
IN['nus'] = calc(wi, r, 'Poisson ratio of sand (drained)', 'silt/clay undrained ν = 0.5', 'ν =', 0.3, '-', kind='in'); r += 1
IN['rcr'] = calc(wi, r, 'Recompression ratio Cr / Cc', 'Das (2019)', 'Cr/Cc =', 0.2, '-', kind='in'); r += 1
IN['sup'] = calc(wi, r, "Strength ratio su / σ'p", 'Mesri (1975)', 'ratio =', 0.22, '-', kind='in'); r += 1
IN['Tv'] = calc(wi, r, 'Time factor for 90 % consolidation', 'Terzaghi', 'Tv =', 0.848, '-', '0.000', 'in'); r += 2
for s in ["Compression index Cc = 0.009 (LL - 10) (Terzaghi & Peck, 1967).",
          "Sand friction angle φ' = 27.1 + 0.3 N60 - 0.00054 N60² (Wolff, 1989).",
          "Silt/clay friction angle sin φ' = 0.8 - 0.094 ln(PI) (Terzaghi, Peck & Mesri, 1996), limited to 20° - 30°.",
          "No oedometer test is available; Cc, Cr, σ'p and Cv are correlated values."]:
    text(wi, r, s, f_note); r += 1
r += 1
text(wi, r, '3. COEFFICIENT OF CONSOLIDATION vs LIQUID LIMIT (NAVFAC DM-7.1)', f_b); r += 1
sp = table_header(wi, r, [('LL', '%', 6), ('Cv', 'cm²/s', 8), ('log10 Cv', '-', 8)]); r += 2
CV = [(30, 3e-3), (40, 1.5e-3), (50, 7e-4), (60, 4e-4), (80, 2e-4), (100, 1.2e-4), (120, 8e-5)]
cv0 = r
for k, (ll, cv) in enumerate(CV):
    table_row(wi, r, sp, [ll, cv, f'=LOG10(J{r})'], ['0', '0.0E+00', '0.000'], inputs=(0, 1), last=k == len(CV) - 1); r += 1
LLT, LCV = f'Input!$C${cv0}:$C${r - 1}', f'Input!$R${cv0}:$R${r - 1}'
# fix: log column references the Cv cell (2nd span starts col 9 = I)
for k in range(len(CV)):
    rr = cv0 + k
    wi.cell(rr, sp[2][0]).value = f'=LOG10({CL(sp[1][0])}{rr})'
LLT = f'Input!${CL(sp[0][0])}${cv0}:${CL(sp[0][0])}${cv0 + len(CV) - 1}'
LCV = f'Input!${CL(sp[2][0])}${cv0}:${CL(sp[2][0])}${cv0 + len(CV) - 1}'
P.close()
IQ = {k: 'Input!' + v for k, v in IN.items()}

# =====================================================================================
HELP = ['top', 'bot', 'h', 'z', 'typ', 'N', 'g', 'sv', 's0', 'LL', 'PL', 'PI', 'e0', 'cu', 'phi', 'E', 'nu', 'm', 'n', 'I', 'ds',
        'F1', 'F2', 'keep', 'Si', 'Cc', 'Cr', 'sp', 'Sc', 'Cv', 't90', 'hb0', 'hbr', 'St', 'ket']
HC = {k: CL(35 + i) for i, k in enumerate(HELP)}          # helper calculation columns AI.. (outside print area)
summary = {}
for bh in BHS:
    ws = wb.create_sheet(bh)
    setup(ws)
    for i, k in enumerate(HELP):
        ws.column_dimensions[HC[k]].width = 9 if k != 'ket' else 40
    g_un, g_sat = B.GAM[bh]
    sub = B.merge(B.sublayers(bh), 41)
    fd = B.fine_default(bh)
    n = len(sub)
    P = Page(ws, 'BEARING CAPACITY AND SETTLEMENT OF SUBGRADE', f'{bh} ({B.AREA[bh]})')
    q, Bw, Lw, HB = IQ['q'], IQ['B'], IQ['L'], IQ['Hb']

    # ---------------- PAGE 1 : general + soil parameters ----------------
    r = P.new(r if 'r' in dir() else None)
    text(ws, r, '1. GENERAL', f_b); r += 1
    text(ws, r, 'This sheet evaluates the bearing capacity, lateral squeeze, immediate settlement and primary consolidation settlement', f_n); r += 1
    text(ws, r, 'of the subgrade below the wellpad embankment (2 m fill + 100 kPa uniform surcharge) at the borehole location.', f_n); r += 2
    text(ws, r, '1.1. Load and loaded area', f_b); r += 1
    calc(ws, r, 'Height of embankment fill', '(Input)', 'Hf =', f'={IQ["H"]}', 'm'); r += 1
    calc(ws, r, 'Uniform surcharge', '(Input)', 'qs =', f'={IQ["qs"]}', 'kPa', '0.0'); r += 1
    calc(ws, r, 'Total pressure on subgrade', 'q = Hf·γf + qs', 'q =', f'={q}', 'kPa', '0.0'); r += 1
    calc(ws, r, 'Loaded area (assumed)', 'B × L', 'B =', f'={Bw}', 'm', '0.0'); r += 1
    calc(ws, r, '', '', 'L =', f'={Lw}', 'm', '0.0'); r += 1
    calc(ws, r, 'Embankment side slope', '1V : nH', 'n =', f'={IQ["slope"]}', '-', '0.0'); r += 2
    text(ws, r, '1.2. Borehole data', f_b); r += 1
    GW = calc(ws, r, 'Design groundwater level (below ground)', 'liquefaction report', 'GWL =', B.GWL[bh], 'm', '0.0', 'in'); r += 1
    G1 = calc(ws, r, 'Unit weight above GWL', 'report Table 5.1', 'γ =', g_un, 'kN/m³', '0.0', 'in'); r += 1
    G2 = calc(ws, r, 'Unit weight below GWL', 'report Table 5.1', 'γsat =', g_sat, 'kN/m³', '0.0', 'in'); r += 1
    DR = calc(ws, r, 'Replacement depth (design condition)', '0 = no improvement', 'Dr =', {'BH-01': 1.0, 'BH-04': 1.0}.get(bh, 0.0), 'm', '0.0', 'in'); r += 2
    r = figure(ws, r, 'fig1_section', 'Figure 1. Typical cross-section of the wellpad embankment, design load and subsoil')
    r = P.new(r)
    text(ws, r, '2. SOIL PARAMETERS', f_b); r += 1
    text(ws, r, f'Table 1. Soil parameters of {bh} (sub-layers = SPT intervals)', f_n); r += 1
    t1 = table_header(ws, r, [('No', '', 2), ('Top', 'm', 3), ('Bottom', 'm', 3), ('Soil', '', 3), ('N-SPT', '', 3), ('γ', 'kN/m³', 3),
                              ('LL', '%', 2), ('PI', '%', 2), ('e0', '-', 2), ('cu', 'kPa', 3), ("φ'", '°', 3)]); r += 2
    R0 = r; RN = R0 + n - 1
    for i, s in enumerate(sub):
        rr = R0 + i
        h = lambda k: f'{HC[k]}{rr}'
        p = lambda k: f'{HC[k]}{rr - 1}'
        ls = B.labstats(bh, s['top'], s['bot'])
        fine = s['cat'] != 'Sand'
        note = []
        if s['src'] != bh:
            note.append(f'data copied from {s["src"]}')
        if i == 0:
            note.append('0-1.5 m not tested: N of next layer')
        if fine and ls['LL'] is None:
            note.append('LL/PL/e0 = borehole average')
        LL = (ls['LL'] or fd['LL']) if fine else None
        PL = (ls['PL'] or fd['PL']) if fine else None
        e0 = ls['e'] or (fd['e'] if fine else 0.65)
        sand = f'{h("typ")}="Sand"'
        M = f'({Lw}/{Bw})'; Nn = f'({h("bot")}/({Bw}/2))'
        s1, s2, s3 = f'SQRT({M}^2+{Nn}^2)', f'SQRT({M}^2+{Nn}^2+1)', f'SQRT({M}^2+1)'
        F1 = f'({M}*LN((1+{s3})*{s1}/({M}*(1+{s2})))+LN(({M}+{s3})*SQRT(1+{Nn}^2)/({M}+{s2})))/PI()'
        F2 = f'{Nn}/(2*PI())*ATAN({M}/({Nn}*{s2}))'
        F1t = '0' if i == 0 else p('F1'); F2t = '0' if i == 0 else p('F2')
        LLi = f'MIN(MAX({h("LL")},30),120)'
        idx = f'MIN(MATCH({LLi},{LLT},1),6)'
        V = {
            'top': s['top'], 'bot': s['bot'], 'h': f'={h("bot")}-{h("top")}', 'z': f'=({h("top")}+{h("bot")})/2', 'typ': s['cat'], 'N': s['N'],
            'g': f'=IF({h("z")}>{GW},{G2},{G1})',
            'sv': f'={h("g")}*{h("h")}' if i == 0 else f'={p("sv")}+{h("g")}*{h("h")}',
            's0': f'={h("sv")}-{h("g")}*{h("h")}/2-10*MAX({h("z")}-{GW},0)',
            'LL': LL, 'PL': PL, 'PI': f'=IF({h("LL")}="","",{h("LL")}-{h("PL")})', 'e0': round(e0, 3),
            'cu': f'=IF({sand},0,{IQ["kcu"]}*{h("N")})',
            'phi': (f'=IF({sand},27.1+0.3*MIN({IQ["ce"]}*{h("N")},60)-0.00054*MIN({IQ["ce"]}*{h("N")},60)^2,'
                    f'MIN(MAX(DEGREES(ASIN(0.8-0.094*LN(MAX(IF({h("PI")}="",10,{h("PI")}),5)))),20),30))'),
            'E': f'=IF({sand},{IQ["kEs"]}*{IQ["ce"]}*{h("N")},{IQ["kEu"]}*{h("cu")})',
            'nu': f'=IF({sand},{IQ["nus"]},0.5)',
            'm': f'={Bw}/2/{h("z")}', 'n': f'={Lw}/2/{h("z")}',
            'I': (f'=1/(4*PI())*(2*{h("m")}*{h("n")}*SQRT({h("m")}^2+{h("n")}^2+1)/({h("m")}^2+{h("n")}^2+1+{h("m")}^2*{h("n")}^2)'
                  f'*({h("m")}^2+{h("n")}^2+2)/({h("m")}^2+{h("n")}^2+1)'
                  f'+ATAN2({h("m")}^2+{h("n")}^2+1-{h("m")}^2*{h("n")}^2,2*{h("m")}*{h("n")}*SQRT({h("m")}^2+{h("n")}^2+1)))'),
            'ds': f'=4*{q}*{h("I")}',
            'F1': '=' + F1, 'F2': '=' + F2,
            'keep': f'=MAX(0,{h("bot")}-MAX({h("top")},{DR}))/{h("h")}*IF({h("ds")}>={IQ["rinf"]}*{h("s0")},1,0)',
            'Si': (f'=4*{q}*({Bw}/2)*(1-{h("nu")}^2)/{h("E")}*(({h("F1")}-{F1t})+(1-2*{h("nu")})/(1-{h("nu")})*({h("F2")}-{F2t}))*1000*{h("keep")}'),
            'Cc': f'=IF({sand},0,0.009*({h("LL")}-10))',
            'Cr': f'={h("Cc")}*{IQ["rcr"]}',
            'sp': f'=IF({sand},0,MAX({h("cu")}/{IQ["sup"]},{h("s0")}))',
            'Sc': (f'=IF({sand},0,1000*{h("h")}/(1+{h("e0")})*IF({h("s0")}+{h("ds")}<={h("sp")},{h("Cr")}*LOG10(({h("s0")}+{h("ds")})/{h("s0")}),'
                   f'{h("Cr")}*LOG10({h("sp")}/{h("s0")})+{h("Cc")}*LOG10(({h("s0")}+{h("ds")})/{h("sp")})))*{h("keep")}'),
            'Cv': (f'=IF({sand},0,10^(INDEX({LCV},{idx})+({LLi}-INDEX({LLT},{idx}))/(INDEX({LLT},{idx}+1)-INDEX({LLT},{idx}))'
                   f'*(INDEX({LCV},{idx}+1)-INDEX({LCV},{idx}))))'),
            't90': f'=IF(OR({sand},{h("keep")}=0),0,{IQ["Tv"]}*({h("h")}/2)^2/({h("Cv")}*1E-4*31536000))',
            'hb0': f'=MAX(0,MIN({h("bot")},{HB})-{h("top")})',
            'hbr': f'=MAX(0,MIN({h("bot")},{HB})-MAX({h("top")},{DR}))',
            'St': f'={h("Si")}+{h("Sc")}',
            'ket': '; '.join(note),
        }
        for k in HELP:
            c = ws[f'{HC[k]}{rr}']; c.value = V[k]; c.font = Font(name=FN, size=9)
            if k in ('top', 'bot', 'typ', 'N', 'LL', 'PL', 'e0'):
                c.fill = FILL_IN
        table_row(ws, rr, t1, [i + 1, f'={h("top")}', f'={h("bot")}', f'={h("typ")}', f'={h("N")}', f'={h("g")}',
                               f'=IF({h("LL")}="","-",{h("LL")})', f'=IF({h("PI")}="","-",{h("PI")})', f'={h("e0")}',
                               f'={h("cu")}', f'={h("phi")}'],
                  ['0', '0.00', '0.00', '@', '0', '0.0', '0', '0', '0.00', '0;-0;"-"', '0.0'], last=i == n - 1)
    hd_r = R0 - 3
    for k in HELP:
        c = ws[f'{HC[k]}{hd_r + 1}']; c.value = k; c.font = Font(name=FN, size=9, bold=True); c.fill = FILL_HD
    ws[f'{HC["top"]}{hd_r}'] = 'Helper calculation per sub-layer (not printed)'; ws[f'{HC["top"]}{hd_r}'].font = f_b
    r = RN + 1
    text(ws, r, 'cu = k·N (silt/clay); φ\' sand from Wolff (1989), silt/clay from PI; γ from laboratory average above/below GWL.', f_note); r += 1
    if bh in B.SRC:
        text(ws, r, f'Below {B.SRC[bh][1]} m the SPT data of {B.SRC[bh][0]} are used ({bh} drilled to {B.SRC[bh][1]} m only).', f_note)
    hr_ = lambda k: f'{HC[k]}{R0}:{HC[k]}{RN}'

    # ---------------- PAGE 2 : bearing capacity ----------------
    r = P.new(r if 'r' in dir() else None)
    text(ws, r, '3. BEARING CAPACITY OF SUBGRADE', f_b); r += 1
    for s in ['The bearing capacity is evaluated with the general bearing capacity equation (Terzaghi 1943, Vesic 1975; SNI 8460:2017).',
              'Undrained (short-term) condition governs where silt/clay exists within 0 - Hb: qu = cu·Nc·sc.',
              "Drained (long-term) condition for granular soil: qu = 0.5·γ'·B·Nγ·sγ·rγ (c' = 0, load at ground surface, q = 0).",
              "Thin soft layer over firm sand below a wide load (B >> H'): squeezing mechanism, Nc* = 4.14 + 0.5·B/H' ≥ 5.14 (Meyerhof 1974).",
              'The governing ultimate capacity is the smaller of the two; qall = qu / SF.']:
        text(ws, r, s); r += 1
    r += 1
    r = figure(ws, r, 'fig2_bearing', 'Figure 2. Bearing capacity mechanisms considered')
    res = {}
    for tag, title, dr in (('0', '3.1. Without ground improvement', '0'), ('r', '3.2. With replacement of soil 0 - Dr (design condition)', DR)):
        text(ws, r, title, f_b); r += 1
        ovl = hr_('hb0') if tag == '0' else hr_('hbr')
        drh = f'MIN({dr},{HB})'
        hf = calc(ws, r, 'Thickness of silt/clay within 0 - Hb', 'Σh (fine soil)', 'Σh =', f'=SUMPRODUCT(--({hr_("typ")}<>"Sand"),{ovl})', 'm'); r += 1
        cu = calc(ws, r, 'Average undrained shear strength', 'weighted by thickness', 'cu =',
                  f'=IF({hf}=0,0,SUMPRODUCT(--({hr_("typ")}<>"Sand"),{ovl},{hr_("cu")})/{hf})', 'kPa', '0.0'); r += 1
        phi = calc(ws, r, 'Average effective friction angle', 'incl. replacement fill', "φ' =",
                   f'=(SUMPRODUCT({ovl},{hr_("phi")})+{drh}*{IQ["phr"]})/{HB}', '°', '0.0'); r += 1
        gam = calc(ws, r, 'Average effective unit weight', '0 - Hb', "γ' =",
                   f'=(SUMPRODUCT({ovl},{hr_("g")})+{drh}*{IQ["gr"]})/{HB}-10*MAX(0,{HB}-{GW})/{HB}', 'kN/m³', '0.0'); r += 1
        ncs = calc(ws, r, "Bearing factor, thin soft layer (H' = Σh)", "Nc* = 4.14 + 0.5·B/H' ≥ 5.14", 'Nc* =',
                   f'=IF({hf}=0,5.14,MAX(5.14,IF({IQ["ncmax"]}>0,MIN({IQ["ncmax"]},4.14+0.5*{Bw}/{hf}),4.14+0.5*{Bw}/{hf})))', '-'); r += 1
        sc = calc(ws, r, 'Shape factor (undrained)', 'sc = 1 + 0.2 B/L', 'sc =', f'=1+0.2*{Bw}/{Lw}', '-'); r += 1
        quu = calc(ws, r, 'Ultimate capacity, undrained', 'qu = cu·Nc*·sc', 'qu,u =', f'=IF({hf}=0,"n/a",{cu}*{ncs}*{sc})', 'kPa', '0.0'); r += 1
        nq = calc(ws, r, 'Bearing capacity factor', 'Nq = e^(π·tanφ)·tan²(45+φ/2)', 'Nq =', f'=EXP(PI()*TAN(RADIANS({phi})))*TAN(RADIANS(45+{phi}/2))^2', '-'); r += 1
        ng = calc(ws, r, 'Bearing capacity factor (Vesic)', 'Nγ = 2(Nq+1)·tanφ', 'Nγ =', f'=2*({nq}+1)*TAN(RADIANS({phi}))', '-'); r += 1
        sg = calc(ws, r, 'Shape and size factors', 'sγ = 1-0.4B/L; rγ = 1-0.25log(B/2)', 'sγ·rγ =',
                  f'=(1-0.4*{Bw}/{Lw})*IF({Bw}>2,1-0.25*LOG10({Bw}/2),1)', '-', '0.000'); r += 1
        qud = calc(ws, r, 'Ultimate capacity, drained', "qu = 0.5·γ'·B·Nγ·sγ·rγ", 'qu,d =', f'=0.5*{gam}*{Bw}*{ng}*{sg}', 'kPa', '0.0'); r += 1
        qu = calc(ws, r, 'Governing ultimate bearing capacity', 'min (undrained, drained)', 'qu =', f'=IF({hf}=0,{qud},MIN({quu},{qud}))', 'kPa', '0.0', 'res'); r += 1
        qa = calc(ws, r, 'Allowable bearing capacity', 'qall = qu / SF', 'qall =', f'={qu}/{IQ["SF"]}', 'kPa', '0.0', 'res'); r += 1
        calc(ws, r, 'Working pressure', 'q = Hf·γf + qs', 'q =', f'={q}', 'kPa', '0.0'); r += 1
        fs = calc(ws, r, 'Actual factor of safety', 'FS = qu / q', 'FS =', f'={qu}/{q}', '-'); r += 1
        hall = calc(ws, r, 'Allowable fill height with surcharge', 'Hall = (qall - qs) / γf', 'Hall =', f'=MAX(0,({qa}-{IQ["qs"]})/{IQ["gf"]})', 'm'); r += 1
        W(ws, r, 3, 14, 'Bearing capacity check', f_b)
        chk = W(ws, r, 24, 31, f'=IF({qa}>={q},"qall ≥ q …. OK!!","qall < q …. NG!!")', f_ok, align=C_C,
                border=Border(left=th, right=th, top=th, bottom=th))
        chk = f'$X${r}'; r += 2
        res[tag] = dict(hf=hf, cu=cu, phi=phi, qa=qa, chk=chk, fs=fs, hall=hall, nc=ncs)
    r = P.new(r)
    text(ws, r, '3.3. Granular soil - allowable pressure for 25 mm settlement (Meyerhof 1965, Bowles 1996)', f_b); r += 1
    nb = calc(ws, r, 'Average N60 of sand layers within 0 - Hb', '', 'N60 =',
              f'=IFERROR(SUMPRODUCT(--({hr_("typ")}="Sand"),{hr_("hb0")},{hr_("N")})/SUMPRODUCT(--({hr_("typ")}="Sand"),{hr_("hb0")})*{IQ["ce"]},"n/a")', '-', '0.0'); r += 1
    qam = calc(ws, r, 'Allowable pressure (S = 25 mm, Kd = 1)', 'qa = (N55/0.08)·((B+0.3)/B)²', 'qa =',
               f'=IFERROR(({nb}*60/55)/0.08*(({Bw}+0.3)/{Bw})^2,"n/a")', 'kPa', '0.0'); r += 1
    text(ws, r, 'For granular soil the shear capacity is normally far above the applied load; settlement governs the design (Section 4).', f_note); r += 2
    text(ws, r, '3.4. Lateral squeeze of soft layer at the embankment toe (Silvestri 1983; FHWA-NHI-06-088)', f_b); r += 1
    for s in ['A thin soft layer below the embankment edge can be squeezed laterally towards the toe (toe heave, edge settlement,',
              'lateral load on piles and casings near the platform edge). Surcharge qs is conservatively assumed to act up to the crest.',
              "      FS = 2·cu / (γf·Ds·tanθ) + 4.14·cu / q_edge ,   q_edge = γf·Hf + qs ,   valid for Ds < b = n·Hf"]:
        text(ws, r, s, f_i if s.startswith('      ') else f_n); r += 1
    r += 1
    r = figure(ws, r, 'fig3_squeeze', 'Figure 3. Lateral squeeze of a soft layer at the embankment toe')
    tth = calc(ws, r, 'Slope inclination', 'tanθ = 1/n', 'tanθ =', f'=1/{IQ["slope"]}', '-', '0.000'); r += 1
    bsl = calc(ws, r, 'Horizontal length of side slope', 'b = n·Hf', 'b =', f'={IQ["slope"]}*{IQ["H"]}', 'm'); r += 1
    sq = {}
    for tag, title in (('0', 'without improvement'), ('r', 'design condition (replacement Dr)')):
        ds_ = res[tag]['hf']; cu_ = res[tag]['cu']
        calc(ws, r, f'Soft layer thickness - {title}', 'Ds = Σh (Section 3.1/3.2)', 'Ds =', f'={ds_}', 'm'); r += 1
        f_ = calc(ws, r, f'FS against lateral squeeze - {title}', '', 'FS =',
                  f'=IF({ds_}=0,"n/a",2*{cu_}/({IQ["gf"]}*{ds_}*{tth})+4.14*{cu_}/{q})', '-', '0.00', 'res'); r += 1
        W(ws, r, 3, 14, '   check')
        W(ws, r, 24, 31, f'=IF({ds_}=0,"No soft layer - n/a",IF({f_}>={IQ["FSsq"]},"FS ≥ FS req …. OK!!","FS < FS req …. NG!!"))', f_ok, align=C_C,
          border=Border(left=th, right=th, top=th, bottom=th))
        sq[tag] = dict(fs=f_, chk=f'$X${r}'); r += 1
    text(ws, r, 'Where Ds ≥ b the toe mechanism is a general bearing failure of the slope and the check reverts to Sections 3.1/3.2.', f_note); r += 1

    # ---------------- PAGE 3 : immediate settlement ----------------
    r = P.new(r if 'r' in dir() else None)
    text(ws, r, '4. IMMEDIATE SETTLEMENT (design condition)', f_b); r += 1
    for s in ['Layered elastic method (Steinbrenner 1934; Bowles 1996) at the centre of a flexible B × L area:',
              '      Si = 4·q·(B/2)·(1 - ν²)/E · [ΔF1 + (1 - 2ν)/(1 - ν)·ΔF2]   for each sub-layer.',
              'Sand: drained modulus E = k·N60 and ν of sand; silt/clay: undrained modulus Eu = k·cu and ν = 0.5.',
              "Layers below the influence depth (Δσz < ratio·σ'v0, Input) and soil replaced to depth Dr are excluded."]:
        text(ws, r, s); r += 1
    text(ws, r, f'Table 2. Immediate settlement of {bh}', f_n); r += 1
    t2 = table_header(ws, r, [('No', '', 2), ('Top', 'm', 3), ('Bottom', 'm', 3), ('Soil', '', 3), ('E', 'kPa', 3), ('ν', '-', 2),
                              ("σ'v0", 'kPa', 3), ('Δσz', 'kPa', 3), ('F1', '-', 2), ('F2', '-', 2), ('Si', 'mm', 3)]); r += 2
    for i in range(n):
        rr = R0 + i; h = lambda k: f'{HC[k]}{rr}'
        table_row(ws, r, t2, [i + 1, f'={h("top")}', f'={h("bot")}', f'={h("typ")}', f'={h("E")}', f'={h("nu")}', f'={h("s0")}',
                              f'={h("ds")}', f'={h("F1")}', f'={h("F2")}', f'={h("Si")}'],
                  ['0', '0.00', '0.00', '@', '#,##0', '0.0', '0.0', '0.0', '0.000', '0.000', '0.0'], last=i == n - 1); r += 1
    r += 1
    si = calc(ws, r, 'Total immediate settlement', 'Si = Σ Si', 'Si =', f'=SUM({hr_("Si")})', 'mm', '0.0', 'res'); r += 1
    sis = calc(ws, r, '   - contribution of sand layers', '', '', f'=SUMPRODUCT(--({hr_("typ")}="Sand"),{hr_("Si")})', 'mm', '0.0'); r += 1
    calc(ws, r, '   - contribution of silt/clay layers', '', '', f'={si}-{sis}', 'mm', '0.0'); r += 1

    # ---------------- PAGE 4 : stress figure + consolidation method ----------------
    r = P.new(r if 'r' in dir() else None)
    r = figure(ws, r, 'fig4_stress', 'Figure 4. Stress distribution below the loaded area and layered elastic settlement')
    text(ws, r, '5. PRIMARY CONSOLIDATION SETTLEMENT (silt/clay layers, design condition)', f_b); r += 1
    for s in ["Sc = h/(1+e0) · [Cr·log(σ'p/σ'v0) + Cc·log((σ'v0+Δσz)/σ'p)]  (Das 2019); Δσz from Boussinesq at the centre of the area;",
              "σ'p = max(cu/0.22 ; σ'v0) (Mesri 1975); t90 = Tv·Hdr²/Cv with double drainage (Hdr = h/2), layers bounded by sand."]:
        text(ws, r, s); r += 1
    r += 1
    r = figure(ws, r, 'fig5_elog', 'Figure 5. Compression curve used for the consolidation settlement', width=800)
    r = P.new(r)
    text(ws, r, f'Table 3. Consolidation settlement of {bh}', f_n); r += 1
    t3 = table_header(ws, r, [('No', '', 2), ('Top', 'm', 3), ('Soil', '', 3), ("σ'v0", 'kPa', 3), ('Δσz', 'kPa', 3), ("σ'p", 'kPa', 3),
                              ('e0', '-', 2), ('Cc', '-', 2), ('Cr', '-', 2), ('Sc', 'mm', 3), ('t90', 'year', 3)]); r += 2
    for i in range(n):
        rr = R0 + i; h = lambda k: f'{HC[k]}{rr}'
        table_row(ws, r, t3, [i + 1, f'={h("top")}', f'={h("typ")}', f'={h("s0")}', f'={h("ds")}', f'={h("sp")}', f'={h("e0")}',
                              f'={h("Cc")}', f'={h("Cr")}', f'={h("Sc")}', f'={h("t90")}'],
                  ['0', '0.00', '@', '0.0', '0.0', '0;-0;"-"', '0.00', '0.000;-0;"-"', '0.000;-0;"-"', '0.0;-0;"-"', '0.00;-0;"-"'],
                  last=i == n - 1); r += 1
    r += 1
    scs = calc(ws, r, 'Total consolidation settlement', 'Sc = Σ Sc', 'Sc =', f'=SUM({hr_("Sc")})', 'mm', '0.0', 'res'); r += 1
    t90 = calc(ws, r, 'Time for 90 % consolidation (slowest layer)', 't90 = max', 't90 =', f'=MAX({hr_("t90")})', 'year', '0.00'); r += 1

    # ---------------- PAGE 5 : summary ----------------
    r = P.new(r if 'r' in dir() else None)
    text(ws, r, '6. SUMMARY', f_b); r += 2
    calc(ws, r, 'Allowable bearing capacity - without improvement', '', 'qall =', f'={res["0"]["qa"]}', 'kPa', '0.0'); r += 1
    W(ws, r, 3, 14, '   check'); W(ws, r, 24, 31, f'={res["0"]["chk"]}', f_ok, align=C_C); r += 1
    calc(ws, r, 'Allowable bearing capacity - design condition', 'Dr from Section 1.2', 'qall =', f'={res["r"]["qa"]}', 'kPa', '0.0'); r += 1
    W(ws, r, 3, 14, '   check'); W(ws, r, 24, 31, f'={res["r"]["chk"]}', f_ok, align=C_C); r += 1
    calc(ws, r, 'FS against lateral squeeze - design condition', 'Section 3.4', 'FS =', f'={sq["r"]["fs"]}', '-', '0.00'); r += 1
    W(ws, r, 3, 14, '   check'); W(ws, r, 24, 31, f'={sq["r"]["chk"]}', f_ok, align=C_C); r += 1
    calc(ws, r, 'Immediate settlement (layered elastic)', 'Section 4', 'Si =', f'={si}', 'mm', '0.0'); r += 1
    calc(ws, r, 'Consolidation settlement', 'Section 5', 'Sc =', f'={scs}', 'mm', '0.0'); r += 1
    calc(ws, r, 'Time for 90 % consolidation', 'Section 5', 't90 =', f'={t90}', 'year', '0.00'); r += 1
    stot = calc(ws, r, 'Total long-term settlement', 'S = Si + Sc', 'S =', f'={si}+{scs}', 'mm', '0.0', 'res'); r += 1
    W(ws, r, 3, 14, 'Settlement check', f_b)
    W(ws, r, 24, 31, f'=IF({stot}<={IQ["Slim"]},"S ≤ S all …. OK!!","S > S all …. NG!!")', f_ok, align=C_C,
      border=Border(left=th, right=th, top=th, bottom=th)); r += 2
    text(ws, r, 'Post-liquefaction settlement during the design earthquake (liquefaction report, Table 5.4: 100 - 383 mm) is not included', f_note); r += 1
    text(ws, r, 'and shall be assessed separately.', f_note); r += 2
    text(ws, r, f'Figure 6. Settlement per sub-layer - {bh}', f_b); r += 1
    ch = ScatterChart(); ch.style = 13; ch.title = None
    ch.y_axis.title = 'Depth (m)'; ch.x_axis.title = 'Settlement per sub-layer (mm)'; ch.y_axis.scaling.orientation = 'maxMin'
    ch.x_axis.delete = False; ch.y_axis.delete = False; ch.legend.position = 'b'
    ch.scatterStyle = 'lineMarker'
    for k, nm, col, sym, wid in (('Si', 'Immediate Si', '0072B2', 'circle', 19050), ('Sc', 'Consolidation Sc', 'E69F00', 'square', 19050),
                                 ('St', 'Total S = Si + Sc', 'C00000', 'triangle', 28575)):
        ci = 35 + HELP.index(k)
        srs = Series(Reference(ws, min_col=35 + HELP.index('z'), min_row=R0, max_row=RN), Reference(ws, min_col=ci, min_row=R0, max_row=RN), title=nm)
        srs.graphicalProperties.line.solidFill = col; srs.graphicalProperties.line.width = wid
        srs.marker.symbol = sym; srs.marker.size = 6
        srs.marker.graphicalProperties.solidFill = col; srs.marker.graphicalProperties.line.solidFill = col
        srs.smooth = False
        ch.series.append(srs)
    ch.width = 17; ch.height = 11
    ws.add_chart(ch, f'D{r}')
    P.close()
    summary[bh] = dict(res=res, sq=sq, dr=DR, qam=qam, si=si, sc=scs, t90=t90, st=stot)

# =====================================================================================
ws = wb.create_sheet('Summary', 1)
setup(ws)
P = Page(ws, 'SUMMARY - BEARING CAPACITY AND SETTLEMENT', 'All boreholes (BH-01 to BH-05)')
r = P.new(r if 'r' in dir() else None)
text(ws, r, 'SUMMARY OF BEARING CAPACITY AND SETTLEMENT OF SUBGRADE', f_b); r += 1
W(ws, r, 3, 31, '="Load: "&TEXT(Input!' + IN['H'] + ',"0.0")&" m embankment fill + "&TEXT(Input!' + IN['qs'] + ',"0")&" kPa uniform surcharge  →  q = "&TEXT(Input!' + IN['q'] + ',"0")&" kPa;  loaded area "&TEXT(Input!' + IN['B'] + ',"0")&" m × "&TEXT(Input!' + IN['L'] + ',"0")&" m"', f_n); r += 2
cols = [('Item', '', 10), ('Unit', '', 3)] + [(bh, B.AREA[bh], 3) for bh in BHS]
sp = table_header(ws, r, cols); r += 2
items = [('Silt/clay thickness within 0 - Hb', 'm', lambda s: s['res']['0']['hf'], '0.0'),
         ('Average cu (silt/clay)', 'kPa', lambda s: s['res']['0']['cu'], '0.0'),
         ("Average φ'", '°', lambda s: s['res']['0']['phi'], '0.0'),
         ('qall - without improvement', 'kPa', lambda s: s['res']['0']['qa'], '#,##0'),
         ('FS - without improvement', '-', lambda s: s['res']['0']['fs'], '0.00'),
         ('Check - without improvement', '', lambda s: s['res']['0']['chk'], '@'),
         ('Replacement depth Dr', 'm', lambda s: s['dr'], '0.0'),
         ('qall - design condition', 'kPa', lambda s: s['res']['r']['qa'], '#,##0'),
         ('FS - design condition', '-', lambda s: s['res']['r']['fs'], '0.00'),
         ('Check - design condition', '', lambda s: s['res']['r']['chk'], '@'),
         ("Nc* (thin soft layer) - design", '-', lambda s: s['res']['r']['nc'], '0.00'),
         ('FS lateral squeeze - design', '-', lambda s: s['sq']['r']['fs'], '0.00'),
         ('Check lateral squeeze - design', '', lambda s: s['sq']['r']['chk'], '@'),
         ('Allowable fill height (design)', 'm', lambda s: s['res']['r']['hall'], '0.0'),
         ('qa for 25 mm (granular soil)', 'kPa', lambda s: s['qam'], '0'),
         ('Immediate settlement Si', 'mm', lambda s: s['si'], '0.0'),
         ('Consolidation settlement Sc', 'mm', lambda s: s['sc'], '0.0'),
         ('Time for 90 % consolidation', 'year', lambda s: s['t90'], '0.00'),
         ('Total settlement S = Si + Sc', 'mm', lambda s: s['st'], '0.0')]
for k, (lab, u, f, fm) in enumerate(items):
    vals = [lab, u] + [f"='{bh}'!{f(summary[bh]).replace('$X$', '$X$')}" for bh in BHS]
    table_row(ws, r, sp, vals, ['@', '@'] + [fm] * 5, last=k == len(items) - 1)
    ws.cell(r, sp[0][0]).alignment = C_L
    if 'Check' in lab:
        for c1, c2 in sp[2:]:
            ws.cell(r, c1).font = Font(name=FN, size=9, bold=True, color='002060')
    r += 1
r += 1
notes = ['NOTES',
         '1. Bearing capacity: undrained (cu·Nc) where silt/clay exists within 0 - Hb; drained (Vesic) for granular soil. For a wide loaded area the',
         '    Nγ term gives a very large drained capacity, i.e. shear failure does not govern; settlement governs the design.',
         "2. Granular soil: (a) drained bearing capacity with φ' from N-SPT; (b) allowable pressure for 25 mm settlement (Meyerhof/Bowles);",
         '    (c) sand is assumed to undergo immediate (elastic) settlement only - no primary consolidation; (d) loose saturated sand: add post-liquefaction settlement for the seismic case.',
         '3. Loaded area B × L = 74 m × 202 m (design brief). For such a wide area the stress increase reaches great depth, so deeper layers contribute to settlement down to the influence depth (Δσz ≥ 0.2·σ\'v0).',
         '4. Post-liquefaction settlement (liquefaction report, Table 5.4): BH-01 289, BH-02 383, BH-03 227, BH-04 159, BH-05 100 mm - assessed separately.',
         '5. BH-01 at 1.5 m: SPT log = clay N 5, liquefaction report/laboratory = sand SP N 12. The SPT log is used (conservative).',
         '6. BH-01 and BH-04: soft silt/clay 0 - 3 m over dense sand is checked with the squeezing factor Nc* (Meyerhof 1974, optional upper limit in Input) and for lateral squeeze at the toe. Design condition: replacement Dr = 1.0 m with compacted granular fill (input per borehole; all checks respond to Dr).',
         '7. Total settlement exceeds the assumed allowable 100 mm at most boreholes: confirm the settlement criterion; consider preloading/surcharging before construction of settlement-sensitive structures.']
import textwrap
merged = [notes[0]]
for s_ in notes[1:]:
    if s_.startswith('    ') and merged:
        merged[-1] += ' ' + s_.strip()
    else:
        merged.append(s_)
for k, s_ in enumerate(merged):
    for j, ln in enumerate(textwrap.wrap(s_, 112, subsequent_indent='    ') if k else [s_]):
        text(ws, r, ln, f_b if k == 0 else f_note); r += 1
P.close()
wb.calculation.fullCalcOnLoad = True
OUT = 'out/Bearing_Capacity_Settlement_BoraPulu.xlsx'
wb.save(OUT)
print(OUT)
