"""Patch PBCP rev3 workbooks (as edited by user in Excel):
- zona likuefaksi (zL) otomatis dari kolom L/NL
- garis A-A' otomatis dari profil penurunan tanah (εv per lapisan, hanya lapisan L) vs zmax
- downdrag memakai A-A' (bukan zL input) -> tanpa lapisan L, DD = 0
"""
import re, sys, zipfile
import openpyxl
from xlsx_cellwriter import Sheet, N, etree

EPS_DEF = 3.0   # % default εv lapisan L tanpa data


def ev_table():
    wb = openpyxl.load_workbook('liqlat.xlsx', data_only=True)
    out = {}
    for bh in ['BH-01', 'BH-02', 'BH-03', 'BH-04', 'BH-05']:
        ws = wb[bh]; t = {}
        for r in range(90, 104):
            d, n, ev = ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 9).value
            if isinstance(d, (int, float)) and isinstance(ev, (int, float)):
                t[round(d, 2)] = (n, ev * 100)
        out[bh] = t
    return out


def style_of(sh, ref):
    c = sh.cell(ref, create=False)
    return c.get('s') if c is not None else None


def patch(src, dst, own_bh):
    z = zipfile.ZipFile(src)
    I, T, D = (Sheet(z.read(f'xl/worksheets/sheet{i}.xml')) for i in (1, 2, 5))
    wbx = openpyxl.load_workbook(src)
    wt = wbx['Data Tanah']
    r0, rn = 7, 7
    while isinstance(wt.cell(rn + 1, 2).value, (int, float)):
        rn += 1
    rng = lambda c: f"'Data Tanah'!${c}${r0}:${c}${rn}"
    lr = lambda c: f"${c}${r0}:${c}${rn}"

    # ---- Data Tanah: kolom P (εv) & Q (penurunan tanah di atas lapisan) ----
    EV = ev_table()
    order = [own_bh] + [b for b in EV if b != own_bh]
    used = {}
    s_hdr5, s_hdr6 = style_of(T, 'N5'), style_of(T, 'N6')
    T.set('P5', 'εv'); T.set('Q5', 'Penurunan tanah di puncak lapisan')
    T.set('P6', '%'); T.set('Q6', 'mm')
    for ref, s in (('P5', s_hdr5), ('Q5', s_hdr5), ('P6', s_hdr6), ('Q6', s_hdr6)):
        if s: T.cell(ref).set('s', s)
    for r in range(r0, rn + 1):
        top = round(float(wt.cell(r, 2).value), 2); n = wt.cell(r, 7).value
        ev = None
        for b in order:
            v = EV[b].get(top)
            if v and v[0] == n:
                ev = v[1]; used[top] = b; break
        if ev is None and top in EV[own_bh] and wt.cell(r, 8).value == 'L':
            ev = EV[own_bh][top][1]; used[top] = own_bh + '(N beda)'
        sN, sM = style_of(T, f'N{r}'), style_of(T, f'G{r}')
        if ev is not None:
            T.set(f'P{r}', round(ev, 2))
        else:
            T.set(f'P{r}', None)
        T.set_formula(f'Q{r}', f'10*SUMPRODUCT(--({lr("B")}>=B{r}),--({lr("H")}="L"),{lr("D")},'
                               f'{lr("P")}+({lr("P")}="")*Input!$C$24)')
        if sM: T.cell(f'P{r}').set('s', sM)     # kuning = input
        if sN: T.cell(f'Q{r}').set('s', sN)
    T.set(f'A{rn + 4}', "εv = regangan volumetrik pasca-likuefaksi (Ishihara & Yoshimine 1992, tabel penurunan laporan likuefaksi). "
                        "Penurunan tanah di puncak lapisan = Σ εv·tebal lapisan L di bawahnya. Lapisan NL tidak menyumbang penurunan.")
    cols = T.tree.find(N + 'cols')
    if cols is not None:
        for cn in (16, 17):
            etree.SubElement(cols, N + 'col', min=str(cn), max=str(cn), width='11', customWidth='1')

    # ---- Input ----
    I.set('A15', 'Dasar zona likuefaksi (lapisan L terdalam)')
    I.set_formula('C15', f'_xlfn.MAXIFS({rng("C")},{rng("H")},"L")')
    I.set('E15', 'otomatis dari kolom Likuefaksi (L/NL) di sheet Data Tanah; 0 = tidak ada likuefaksi')
    I.set('A20', 'Penurunan tanah di permukaan akibat likuefaksi')
    I.set_formula('C20', f'MAX({rng("Q")})')
    I.set('E20', 'otomatis = Σ εv·tebal lapisan L (Data Tanah kolom P-Q)')
    I.set('A22', "Garis A-A' (batas bawah zona downdrag)")
    I.set('B22', "zA-A'")
    I.set_formula('C22', f'_xlfn.MAXIFS({rng("C")},{rng("Q")},">"&Input!$C$21)')
    I.set('D22', 'm')
    I.set('E22', "otomatis: dasar lapisan terdalam yang penurunan tanahnya > zmax (Caltrans Step 4-5, δpile diabaikan); 0 = tanpa downdrag")
    c22 = I.cell('C22'); c21 = I.cell('C21', create=False)
    if c21 is not None and c21.get('s'): c22.set('s', c21.get('s'))
    I.set('A24', 'εv default lapisan L tanpa data')
    I.set('C24', EPS_DEF); I.set('D24', '%'); I.set('E24', 'dipakai bila kolom εv (Data Tanah) kosong pada lapisan L')
    s_in = style_of(I, 'C23')
    if s_in:
        I.cell('C24').set('s', s_in)
    for ref, src_ref in (('A24', 'A23'), ('D24', 'D23'), ('E24', 'E23'), ('D22', 'D21'), ('B22', 'B21')):
        s = style_of(I, src_ref)
        if s: I.cell(ref).set('s', s)

    # ---- Downdrag: pakai A-A' ----
    n_rep = 0
    for row in D.sd.findall(N + 'row'):
        for c in row.findall(N + 'c'):
            f = c.find(N + 'f')
            if f is None or not f.text or 'Input!$C$15' not in f.text:
                continue
            col = re.match(r'[A-Z]+', c.get('r')).group(0)
            if col == 'E':
                f.text = f.text.replace('Input!$C$15', 'MAX(Input!$C$22,Input!$C$27)')
            else:
                f.text = f.text.replace('Input!$C$15', 'Input!$C$22')
            v = c.find(N + 'v')
            if v is not None: c.remove(v)
            n_rep += 1

    # ---- write package (drop calcChain, force recalc, widen print area) ----
    edited = {'xl/worksheets/sheet1.xml': I.bytes(), 'xl/worksheets/sheet2.xml': T.bytes(), 'xl/worksheets/sheet5.xml': D.bytes()}
    with zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED) as zo:
        for it in z.infolist():
            if it.filename == 'xl/calcChain.xml':
                continue
            data = edited.get(it.filename) or z.read(it.filename)
            if it.filename.startswith('xl/worksheets/sheet') and it.filename.endswith('.xml'):
                # buang nilai cache rumus agar semua dihitung ulang (Excel: fullCalcOnLoad)
                tr = etree.fromstring(data)
                for c in tr.iter(N + 'c'):
                    if c.find(N + 'f') is not None:
                        v = c.find(N + 'v')
                        if v is not None: c.remove(v)
                        c.attrib.pop('t', None)
                data = etree.tostring(tr, xml_declaration=True, encoding='UTF-8', standalone=True)
            if it.filename == 'xl/workbook.xml':
                s = data.decode('utf8')
                s = s.replace("'Data Tanah'!$A$1:$O$", "'Data Tanah'!$A$1:$Q$")
                s = re.sub(r'<calcPr([^>]*?)/>', lambda m: '<calcPr' + re.sub(r'\s*fullCalcOnLoad="[^"]*"', '', m.group(1)) + ' fullCalcOnLoad="1"/>', s, 1)
                data = s.encode()
            elif it.filename == 'xl/_rels/workbook.xml.rels':
                data = re.sub(rb'<Relationship [^>]*calcChain[^>]*/>', b'', data)
            elif it.filename == '[Content_Types].xml':
                data = re.sub(rb'<Override [^>]*calcChain[^>]*/>', b'', data)
            zo.writestr(it, data)
    return n_rep, rn, used


if __name__ == '__main__':
    for bh in ['01', '02', '03', '04', '05']:
        print(bh, patch(f'drv2/PBCP-BH-{bh}_rev3.xlsx', f'out/PBCP-BH-{bh}_rev4.xlsx', f'BH-{bh}'))
