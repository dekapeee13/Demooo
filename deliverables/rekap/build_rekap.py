"""Rekap gaya dalam tiang from the GROUP rerun outputs (16 piles, 2 cellars) -> tidy workbook + SPColumn load pairs.
Usage: python3 build_rekap.py <folder with M1/M3_Wellpad A/D_4x1.gp11o.txt> <out.xlsx>"""
import sys, collections
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

src, out = sys.argv[1], sys.argv[2]
import rerun
DATA = rerun.load(src)
rows = rerun.rows(DATA)
# r = (wellpad, model, lc, pile, z, P, M, V, y, p_soil, stress)
COND = {'M1': 'Non-liquefied', 'M3': 'Liquefied'}
method = lambda lc: 'LRFD' if lc.startswith('LRFD') else 'ASD'

NAVY, GREY = '1F3864', 'F2F2F2'
thin = Side(style='thin', color='BFBFBF')
B = Border(left=thin, right=thin, top=thin, bottom=thin)
HF, HFILL = Font(name='Calibri', size=10, bold=True, color='FFFFFF'), PatternFill('solid', fgColor=NAVY)
F, FB = Font(name='Calibri', size=10), Font(name='Calibri', size=10, bold=True)
RED = Font(name='Calibri', size=10, bold=True, color='C00000')
C = Alignment(horizontal='center', vertical='center', wrap_text=True)

def put(ws, r0, header, data, widths=None, fmt=None, hi=None):
    for j, h in enumerate(header, 1):
        c = ws.cell(r0, j, h); c.font, c.fill, c.alignment, c.border = HF, HFILL, C, B
    ws.row_dimensions[r0].height = 30
    for i, row in enumerate(data, 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(r0 + i, j, v); c.font, c.alignment, c.border = F, C, B
            if i % 2 == 0: c.fill = PatternFill('solid', fgColor=GREY)
            if fmt and fmt[j - 1]: c.number_format = fmt[j - 1]
            if hi and (i - 1, j - 1) in hi: c.font = RED
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(j)].width = max(ws.column_dimensions[get_column_letter(j)].width or 0, w)
    return r0 + len(data) + 2

def title(ws, t, sub=None):
    ws['A1'] = t; ws['A1'].font = Font(name='Calibri', size=13, bold=True, color=NAVY)
    if sub: ws['A2'] = sub; ws['A2'].font = Font(name='Calibri', size=9, italic=True, color='595959')

def setup(ws):
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_view.showGridLines = False

# per (wellpad, model, lc, pile): P (constant), Mmax+z, Vmax+z, head y
pile = {}
for wp, mdl, lc, p, z, P, M, V, y, *_ in rows:
    k = (wp, mdl, lc, int(p)); e = pile.setdefault(k, dict(P=P, M=-1, zM=0, V=-1, zV=0, y=0))
    if abs(M) > e['M']: e['M'], e['zM'] = abs(M), z
    if abs(V) > e['V']: e['V'], e['zV'] = abs(V), z
    if z == 0: e['y'] = abs(y)

wb = openpyxl.Workbook()
# ------------------------------------------------------------------ 1. Summary
ws = wb.active; ws.title = 'Summary'
title(ws, 'PILE INTERNAL FORCES - BORED PILE D1000, GROUP 16 PILES (2 CELLARS, 4x1 PER RIG NODE) WITH CAP INERTIA (+Icap)',
      'P = axial at pile head (+ compression, - tension), constant along the pile. M, V = resultant of both directions. Source: GROUP rerun .gp11o (Drive 02. Output GROUP).')
grp = collections.defaultdict(list)
for (wp, mdl, lc, p), e in pile.items():
    grp[(wp, mdl, method(lc))].append((lc, p, e))
data, hi = [], set()
for (wp, mdl, mt), L in sorted(grp.items()):
    pmax = max(L, key=lambda t: t[2]['P']); pmin = min(L, key=lambda t: t[2]['P'])
    mm = max(L, key=lambda t: t[2]['M']); vm = max(L, key=lambda t: t[2]['V']); ym = max(L, key=lambda t: t[2]['y'])
    data.append([wp, COND[mdl], mt, pmax[2]['P'], f'{pmax[0]} / {pmax[1]}', pmin[2]['P'], f'{pmin[0]} / {pmin[1]}',
                 mm[2]['M'], mm[2]['P'], mm[2]['zM'], f'{mm[0]} / {mm[1]}', vm[2]['V'], vm[2]['zV'], f'{vm[0]} / {vm[1]}', ym[2]['y']])
    if pmin[2]['P'] < 0: hi.add((len(data) - 1, 5))
nxt = put(ws, 4, ['Wellpad', 'Condition', 'Design', 'P max (kN)', 'LC / pile', 'P min (kN)', 'LC / pile', 'M max (kNm)',
                  'P with M max (kN)', 'z M max (m)', 'LC / pile', 'V max (kN)', 'z V max (m)', 'LC / pile', 'Head defl. max (mm)'],
          data, [8, 13, 7, 10, 16, 10, 16, 10, 11, 9, 16, 10, 9, 16, 10],
          [None, None, None, '#,##0', None, '#,##0', None, '#,##0', '#,##0', '0.0', None, '#,##0', '0.0', None, '0.0'], hi)
notes = ['Notes:',
         'Seismic: ASD 0.7E, LRFD 1.0E incl. cap inertia; direction 100/30. M3 (liquefied) run has 6 load cases (ASD1/3/4, LRFD1/3/4).',
         'Red P min = tension in the pile (if any): check tension reinforcement and pull-out capacity.',
         "Cracking moment Mcr = 0.62 sqrt(fc') I/c = 333 kNm (no axial): M max above Mcr -> cracked section, design with SPColumn.",
         'Equilibrium: sum of the 16 pile-head axial forces = vertical load at the cap origin for every load case (checked).',
         'Pile length in the GROUP models (last output depth): ' + ', '.join(f'{k[0]} {k[1]} {v["cases"][0]["piles"][1]["depth"][-1]["x"]:g} m' for k, v in DATA.items())
         + '. Design: WPA 43 m, WPD 48 m.']
for i, t in enumerate(notes):
    ws.cell(nxt + i, 1, t).font = FB if i == 0 else F
setup(ws); ws.freeze_panes = 'D5'

# ------------------------------------------------------------------ 2. SPColumn (LRFD factored pairs)
ws = wb.create_sheet('SPColumn_Pu_Mu')
title(ws, 'FACTORED LOAD PAIRS FOR SPCOLUMN (LRFD) - critical pile per load case',
      'Per load case: pile with P max, pile with P min and pile with M max. Mu = max resultant moment along the pile; Pu = head axial (+ compression).')
data, hi = [], set()
for (wp, mdl, mt), L in sorted(grp.items()):
    if mt != 'LRFD': continue
    for lc in sorted({t[0] for t in L}):
        Lc = [t for t in L if t[0] == lc]
        picks = [('P max', max(Lc, key=lambda t: t[2]['P'])), ('P min', min(Lc, key=lambda t: t[2]['P'])),
                 ('M max', max(Lc, key=lambda t: t[2]['M']))]
        seen = {}
        for crit, (_, p, e) in picks:
            seen.setdefault(p, []).append(crit)
        for p, crits in seen.items():
            e = dict((t[1], t[2]) for t in Lc)[p]
            data.append([wp, COND[mdl], lc, p, ', '.join(crits), e['P'], e['M'], e['zM'], e['V']])
# highlight envelope per wellpad
for wp in ('WPA', 'WPD'):
    idx = [i for i, r in enumerate(data) if r[0] == wp]
    for col, f in ((5, max), (5, min), (6, max)):
        hi.add((f(idx, key=lambda i: data[i][col]), col))
nxt = put(ws, 4, ['Wellpad', 'Condition', 'Load case', 'Pile', 'Critical for', 'Pu (kN)', 'Mu (kNm)', 'z Mu (m)', 'Vu (kN)'],
          data, [8, 13, 16, 6, 18, 10, 10, 9, 10], [None, None, None, '0', None, '#,##0', '#,##0', '0.0', '#,##0'], hi)
ws.cell(nxt, 1, 'Red = governing value per wellpad (P max, P min/tension, M max) - minimum set of points to plot in the SPColumn P-M diagram.').font = F
# compact SPColumn input per wellpad (envelope M1/M3)
r0 = nxt + 2
ws.cell(r0, 1, 'SPColumn factored loads (envelope M1 + M3)').font = Font(name='Calibri', size=11, bold=True, color=NAVY)
inp = []
for wp in ('WPA', 'WPD'):
    sub = [r for r in data if r[0] == wp]
    for lab, r in (('Pu max', max(sub, key=lambda r: r[5])), ('Pu min', min(sub, key=lambda r: r[5])), ('Mu max', max(sub, key=lambda r: r[6]))):
        inp.append([wp, lab, f'{r[1]} - {r[2]} / pile {r[3]}', r[5], r[6]])
put(ws, r0 + 1, ['Wellpad', 'Point', 'Source', 'Pu (kN)', 'Mu (kNm)'], inp, None, [None, None, None, '#,##0', '#,##0'])
setup(ws); ws.freeze_panes = 'A5'


# ------------------------------------------------------------------ 2b. SPColumn input (copy-paste ready)
ws = wb.create_sheet('SPColumn_Input', 2)
title(ws, 'SPCOLUMN FACTORED LOADS - ready to paste (Loads > Factored Loads)',
      'Copy columns B:D (Pu, Mux, Muy) of each block into spColumn. Pu: + compression, - tension (spColumn sign). '
      'Circular section: resultant Mu entered as Mux, Muy = 0. All LRFD load cases x all piles x M1/M3.')
r0, col0 = 4, 1
for wp in ('WPA', 'WPD'):
    lst = sorted(((mdl, lc, p, e) for (w, mdl, lc, p), e in pile.items() if w == wp and method(lc) == 'LRFD'),
                 key=lambda t: (t[0], t[1], t[2]))
    ws.cell(r0, col0, f'{wp} - {len(lst)} load points').font = Font(name='Calibri', size=11, bold=True, color=NAVY)
    dat = [[i, round(e['P'], 1), round(e['M'], 1), 0, f'{COND[mdl]} | {lc} | pile {p} | Mu at z = {e["zM"]:g} m']
           for i, (mdl, lc, p, e) in enumerate(lst, 1)]
    hi = {(max(range(len(dat)), key=lambda i: dat[i][1]), 1), (min(range(len(dat)), key=lambda i: dat[i][1]), 1),
          (max(range(len(dat)), key=lambda i: dat[i][2]), 2)}
    for i, row in enumerate(dat):
        for j, v in enumerate(row):
            c = ws.cell(r0 + 2 + i, col0 + j, v); c.font, c.border = (RED if (i, j) in hi else F), B
            c.alignment = Alignment(horizontal='left' if j == 4 else 'center')
            if j in (1, 2, 3): c.number_format = '0.0'
    for j, h in enumerate(['No.', 'Pu (kN)', 'Mux (kNm)', 'Muy (kNm)', 'Description (not for spColumn)']):
        c = ws.cell(r0 + 1, col0 + j, h); c.font, c.fill, c.alignment, c.border = HF, HFILL, C, B
    for j, w in enumerate([6, 10, 11, 11, 52]):
        ws.column_dimensions[get_column_letter(col0 + j)].width = w
    with open(out.rsplit('/', 1)[0] + f'/SPColumn_Loads_{wp}.csv' if '/' in out else f'SPColumn_Loads_{wp}.csv', 'w') as fh:
        fh.write('\n'.join(f'{r[1]},{r[2]},{r[3]}' for r in dat) + '\n')
    col0 += 7
ws.freeze_panes = 'A6'; setup(ws)

# ------------------------------------------------------------------ 3. Envelope vs depth
ws = wb.create_sheet('Envelope_Depth')
title(ws, 'ENVELOPE OF M, V AND DEFLECTION VS DEPTH (all load cases and piles)')
env = collections.defaultdict(lambda: dict(ML=0, VL=0, MA=0, VA=0, y=0))
for wp, mdl, lc, p, z, P, M, V, y, *_ in rows:
    e = env[(wp, mdl, z)]; s = 'L' if method(lc) == 'LRFD' else 'A'
    e['M' + s] = max(e['M' + s], abs(M)); e['V' + s] = max(e['V' + s], abs(V)); e['y'] = max(e['y'], abs(y))
r0 = 4
for wp in ('WPA', 'WPD'):
    for mdl in ('M1', 'M3'):
        ws.cell(r0, 1, f'{wp} - {mdl} ({COND[mdl]})').font = Font(name='Calibri', size=11, bold=True, color=NAVY)
        zs = sorted(z for (a, b, z) in env if (a, b) == (wp, mdl))
        dat = [[z, env[(wp, mdl, z)]['ML'], env[(wp, mdl, z)]['VL'], env[(wp, mdl, z)]['MA'], env[(wp, mdl, z)]['VA'], env[(wp, mdl, z)]['y']] for z in zs]
        r0 = put(ws, r0 + 1, ['z (m)', 'LRFD M max (kNm)', 'LRFD V max (kN)', 'ASD M max (kNm)', 'ASD V max (kN)', 'Defl. max (mm)'],
                 dat, [9, 12, 12, 12, 12, 12], ['0.0', '#,##0.0', '#,##0.0', '#,##0.0', '#,##0.0', '0.00'])
setup(ws)

# ------------------------------------------------------------------ 4. Raw data
ws = wb.create_sheet('Data')
put(ws, 1, ['Wellpad', 'Model', 'Load case', 'Pile', 'z (m)', 'P (kN)', 'M resultant (kNm)', 'V resultant (kN)', 'Defl. (mm)'],
    [list(r[:9]) for r in rows], [8, 7, 16, 6, 7, 10, 11, 11, 10])
ws.freeze_panes = 'A2'; ws.auto_filter.ref = f'A1:I{len(rows) + 1}'
wb.save(out)

if __name__ == '__main__':
    chk = openpyxl.load_workbook(out)['Summary']
    vals = {(chk.cell(r, 1).value, chk.cell(r, 2).value, chk.cell(r, 3).value): chk.cell(r, 8).value for r in range(5, 13)}
    pm = {(chk.cell(r, 1).value, chk.cell(r, 2).value, chk.cell(r, 3).value): chk.cell(r, 4).value for r in range(5, 13)}
    assert abs(pm[('WPA', 'Non-liquefied', 'LRFD')] - 6880.5) < 0.1, pm   # = Rekap_Wellpad_A, Envelope per Pile
    assert abs(pm[('WPD', 'Non-liquefied', 'ASD')] - 5057.0) < 0.1, pm
    print('ok', out)
