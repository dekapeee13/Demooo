"""Pile-head forces per pile and per combination for pile-cap reinforcement design (GROUP rerun, 16 piles, 2 cellars).
Usage: python3 build_pilecap.py <folder with M1/M3_Wellpad A/D_4x1.gp11o.txt> <out.xlsx>"""
import sys, math
import openpyxl
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as L
import rerun

src, out = sys.argv[1], sys.argv[2]
DATA = rerun.load(src)
XY = rerun.coords(DATA[('WPD', 'M1')]['cases'])
for d in DATA.values():
    assert rerun.coords(d['cases']) == XY            # same layout in all 4 models
NP = len(XY)
COND = {'M1': 'Non-liquefied', 'M3': 'Liquefied'}
ORDER = ['ASD1', 'ASD2(max)', 'ASD2(min)', 'ASD3', 'ASD4', 'LRFD1', 'LRFD2(max)', 'LRFD2(min)', 'LRFD3', 'LRFD4']
CAP_B, CAP_H = 25.0, 12.0
# head[(wp, mdl, lc, pile)] = dict(FX, FY, FZ, MX, MY, MZ, M, V, d)
head = {}
for (wp, mdl), d in DATA.items():
    for c in d['cases']:
        for p, pl in c['piles'].items():
            FX, FY, FZ, MX, MY, MZ = pl['gforce']
            head[(wp, mdl, rerun.lcname(c['name']), p)] = dict(FX=FX, FY=FY, FZ=FZ, MX=MX, MY=MY, MZ=MZ, M=math.hypot(MY, MZ),
                                                             V=math.hypot(FY, FZ), d=math.hypot(*pl['gdisp'][1:3]) * 1000)
lcs = lambda wp, mdl: sorted({k[2] for k in head if k[:2] == (wp, mdl)}, key=ORDER.index)
side = lambda p: 'Y < 0' if XY[p][0] < 0 else 'Y > 0'

NAVY, GREY = '1F3864', 'F2F2F2'
thin = Side(style='thin', color='BFBFBF'); B = Border(left=thin, right=thin, top=thin, bottom=thin)
HF, HFILL = Font(name='Calibri', size=10, bold=True, color='FFFFFF'), PatternFill('solid', fgColor=NAVY)
F, FB = Font(name='Calibri', size=10), Font(name='Calibri', size=10, bold=True)
RED, BLUE = Font(name='Calibri', size=10, bold=True, color='C00000'), Font(name='Calibri', size=10, bold=True, color='0070C0')
C = Alignment(horizontal='center', vertical='center', wrap_text=True); LEFT = Alignment(horizontal='left', vertical='center')
H1 = Font(name='Calibri', size=13, bold=True, color=NAVY); H2 = Font(name='Calibri', size=11, bold=True, color=NAVY)
SUB = Font(name='Calibri', size=9, italic=True, color='595959')

def cell(ws, r, c, v, font=F, fill=None, fmt=None, al=C):
    x = ws.cell(r, c, v); x.font, x.alignment, x.border = font, al, B
    if fill: x.fill = PatternFill('solid', fgColor=fill)
    if fmt: x.number_format = fmt
    return x

def hdr(ws, r, c0, labels, h=30):
    for j, t in enumerate(labels):
        x = ws.cell(r, c0 + j, t); x.font, x.fill, x.alignment, x.border = HF, HFILL, C, B
    ws.row_dimensions[r].height = h

def setup(ws):
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_view.showGridLines = False

wb = openpyxl.Workbook()
# ------------------------------------------------------------------ 1. Layout
ws = wb.active; ws.title = 'Layout'
ws['A1'] = 'PILE IDENTIFICATION AND LAYOUT - GROUP RERUN, 16 BORED PILES D1000 UNDER CAP 25 x 12 m (2 CELLARS)'; ws['A1'].font = H1
ws['A2'] = ('Pile-head coordinates back-calculated from the GROUP output (rigid-cap kinematics: vertical head displacement vs cap '
            'rotation, all load cases, residual < 1e-12 m). Same layout in all 4 runs (WPA/WPD, M1/M3).'); ws['A2'].font = SUB
hdr(ws, 4, 1, ['GROUP pile no.', 'GROUP Y (m)', 'GROUP Z (m)', 'Row', 'Position in row', 'Side of cap',
               'Cap x from left edge (m)', 'Cap y from bottom edge (m)'])
for p in range(1, NP + 1):
    y, z = XY[p]; r = 4 + p
    for j, v in enumerate([p, y, z, 1 if z < 0 else 2, (p - 1) % 8 + 1, side(p), y + CAP_B / 2, z + CAP_H / 2], 1):
        cell(ws, r, j, v, FB if j == 1 else F, GREY if p % 2 == 0 else None, '0.00' if j in (2, 3, 7, 8) else None)
for j, w in enumerate([10, 10, 10, 6, 9, 10, 12, 12], 1): ws.column_dimensions[L(j)].width = w
notes = ['Notes:',
         'GROUP global axes: X = vertical (axial, + compression), Y and Z = horizontal. Origin = cap centre (cap weight/inertia applied at Y = Z = 0).',
         'Rows at Z = -4.75 m (piles 1-8) and Z = +4.75 m (piles 9-16); 8 piles per row at 2.50 m spacing (Y = -8.75 ... +8.75 m).',
         '4x1 under each rig node: piles 1-4 / 9-12 under the cellar at Y < 0, piles 5-8 / 13-16 under the cellar at Y > 0.',
         'Cap x/y assume the GROUP origin at the cap centre and GROUP Y along the 25 m side. Which cellar (and rig nodes 3226/3233/3234/3241)',
         'is at Y < 0 must be confirmed with the GROUP input (.gp11d) - the output does not print the load coordinates.',
         'Edge distance: 12/2 - 4.75 = 1.25 m (cap edge to pile centre) across, 12.5 - 8.75 = 3.75 m along the cap.']
for i, t in enumerate(notes): ws.cell(23 + i, 1, t).font = FB if i == 0 else F
fig, ax = plt.subplots(figsize=(8.2, 4.6))
ax.add_patch(plt.Rectangle((-CAP_B / 2, -CAP_H / 2), CAP_B, CAP_H, fc='#F2F2F2', ec='k', lw=1.4))
for cx in (-5.0, 5.0):
    ax.add_patch(plt.Rectangle((cx - 1.9, -1.75), 3.8, 3.5, fc='white', ec='0.4', lw=1, ls='--'))
    ax.text(cx, 0, 'cellar\n(approx.)', ha='center', va='center', fontsize=7, color='0.4')
for p, (y, z) in XY.items():
    ax.add_patch(plt.Circle((y, z), 0.5, fc='#D9E2F3', ec='#1F3864', lw=1.2))
    ax.text(y, z, str(p), ha='center', va='center', fontsize=9, weight='bold', color='#1F3864')
ax.plot(0, 0, '+', color='#C00000', ms=10); ax.text(0.2, 0.3, 'origin', fontsize=7, color='#C00000')
ax.set_xlim(-13.5, 13.5); ax.set_ylim(-7, 7); ax.set_aspect('equal'); ax.grid(color='0.92')
ax.set_xlabel('GROUP Y (m)'); ax.set_ylabel('GROUP Z (m)')
ax.set_title('Pile numbering (GROUP) - cap 25 x 12 m, 16 piles D1000 at 2.5 m', fontsize=10)
fig.tight_layout(); png = out.rsplit('.', 1)[0] + '_layout.png'; fig.savefig(png, dpi=130); plt.close(fig)
img = XLImage(png); img.anchor = 'J4'; ws.add_image(img)
setup(ws)

# ------------------------------------------------------------------ 2. equilibrium check
ws = wb.create_sheet('Check_Equilibrium')
ws['A1'] = 'EQUILIBRIUM CHECK - sum of pile-head forces vs equivalent load at the cap origin (GROUP Table L)'; ws['A1'].font = H1
hdr(ws, 3, 1, ['Wellpad', 'Model', 'Load combination', 'V at origin (kN)', 'Sum FX piles (kN)', 'Ratio', 'Hy (kN)', 'Sum FY (kN)', 'Hz (kN)', 'Sum FZ (kN)', 'Status'])
r = 4
for (wp, mdl), d in DATA.items():
    for c in d['cases']:
        lc = rerun.lcname(c['name']); V, Hy, Hz = c['load'][:3]
        s = [sum(head[(wp, mdl, lc, p)][k] for p in XY) for k in ('FX', 'FY', 'FZ')]
        for j, v in enumerate([wp, mdl, lc, V, s[0], None, Hy, s[1], Hz, s[2]], 1):
            cell(ws, r, j, v, fmt='#,##0.0' if j >= 4 else None, al=LEFT if j == 3 else C)
        cell(ws, r, 6, f'=E{r}/D{r}', fmt='0.000')
        cell(ws, r, 11, f'=IF(AND(ABS(F{r}-1)<=0.01,ABS(H{r}-G{r})<=0.01*ABS(D{r}),ABS(J{r}-I{r})<=0.01*ABS(D{r})),"OK","CHECK")', FB)
        r += 1
ws.conditional_formatting.add(f'K4:K{r - 1}', CellIsRule(operator='equal', formula=['"CHECK"'], font=RED))
for j, w in enumerate([8, 7, 14, 12, 12, 7, 10, 10, 10, 10, 8], 1): ws.column_dimensions[L(j)].width = w
setup(ws)

# ------------------------------------------------------------------ 3. governing per pile
ws = wb.create_sheet('Governing')
ws['A1'] = 'GOVERNING PILE-HEAD FORCES PER PILE (envelope M1 + M3) - LRFD for cap reinforcement, ASD for service checks'; ws['A1'].font = H1
ws['A2'] = 'M, V = resultant of the head moments (MY, MZ) and shears (FY, FZ). LC = model + load combination.'; ws['A2'].font = SUB
r0 = 4
for wp in ('WPA', 'WPD'):
    ws.cell(r0, 1, f'Wellpad {wp[-1]}').font = H2
    hdr(ws, r0 + 1, 1, ['Pile', 'Y (m)', 'Z (m)', 'Pu max (kN)', 'LC', 'Pu min (kN)', 'LC', 'Mu max (kNm)', 'LC', 'Vu max (kN)', 'LC', 'P max ASD (kN)', 'LC'])
    for p in range(1, NP + 1):
        r = r0 + 1 + p
        ks = [k for k in head if k[0] == wp and k[3] == p]
        lr = [k for k in ks if k[2].startswith('LRFD')]; asd = [k for k in ks if k[2].startswith('ASD')]
        picks = [(max(lr, key=lambda k: head[k]['FX']), 'FX'), (min(lr, key=lambda k: head[k]['FX']), 'FX'),
                 (max(lr, key=lambda k: head[k]['M']), 'M'), (max(lr, key=lambda k: head[k]['V']), 'V'), (max(asd, key=lambda k: head[k]['FX']), 'FX')]
        cell(ws, r, 1, p, FB); cell(ws, r, 2, XY[p][0], fmt='0.00'); cell(ws, r, 3, XY[p][1], fmt='0.00')
        for j, (k, q) in enumerate(picks):
            v = head[k][q]
            cell(ws, r, 4 + 2 * j, v, RED if v < 0 else F, fmt='#,##0.0'); cell(ws, r, 5 + 2 * j, f'{k[1]} {k[2]}')
    r0 += NP + 4
for j, w in enumerate([6, 7, 7, 11, 15, 11, 15, 11, 15, 11, 15, 11, 15], 1): ws.column_dimensions[L(j)].width = w
setup(ws)

# ------------------------------------------------------------------ 4. matrices per wellpad / model
QTY = [('FX', 'Axial force FX (kN)  [+ compression]'), ('FY', 'Shear FY (kN)'), ('FZ', 'Shear FZ (kN)'),
       ('MY', 'Moment MY (kNm)'), ('MZ', 'Moment MZ (kNm)'), ('M', 'Resultant head moment M (kNm)'),
       ('V', 'Resultant head shear V (kN)'), ('d', 'Lateral head deflection (mm)')]
for wp in ('WPA', 'WPD'):
    for mdl in ('M1', 'M3'):
        ws = wb.create_sheet(f'{wp}_{mdl}')
        ws['A1'] = f'PILE-HEAD FORCES - {wp} {mdl} ({COND[mdl]}), GROUP rerun, 16 piles (global GROUP axes)'; ws['A1'].font = H1
        ws['A2'] = 'Forces exerted by the cap on the pile head; reverse the sign for the reaction on the cap. Pile numbers and coordinates: sheet Layout.'; ws['A2'].font = SUB
        L_ = lcs(wp, mdl); r0 = 4
        for q, lab in QTY:
            ws.cell(r0, 1, lab).font = H2
            hdr(ws, r0 + 1, 1, ['Load combination', 'Design'] + [f'P{p}' for p in range(1, NP + 1)] + (['Sum'] if q in ('FX', 'FY', 'FZ') else []), 20)
            for i, lc in enumerate(L_):
                r = r0 + 2 + i; g = GREY if i % 2 else None
                cell(ws, r, 1, lc, FB, g, al=LEFT); cell(ws, r, 2, 'LRFD' if lc.startswith('LRFD') else 'ASD', fill=g)
                for p in range(1, NP + 1):
                    v = head[(wp, mdl, lc, p)][q]
                    cell(ws, r, 2 + p, round(v, 3) if abs(v) > 1e-6 else 0, RED if (q == 'FX' and v < 0) else F, g, '#,##0.0' if q != 'd' else '0.00')
                if q in ('FX', 'FY', 'FZ'):
                    cell(ws, r, 3 + NP, f'=SUM(C{r}:{L(2 + NP)}{r})', FB, g, '#,##0.0')
            r = r0 + 2 + len(L_)
            for k, (nm, fn) in enumerate((('max', 'MAX'), ('min', 'MIN'))):
                cell(ws, r + k, 1, nm, BLUE, 'DDEBF7', al=LEFT); cell(ws, r + k, 2, '', fill='DDEBF7')
                for p in range(1, NP + 1):
                    cl = L(2 + p); cell(ws, r + k, 2 + p, f'={fn}({cl}{r0 + 2}:{cl}{r - 1})', BLUE, 'DDEBF7', '#,##0.0' if q != 'd' else '0.00')
            r0 = r + 4
        ws.column_dimensions['A'].width = 13; ws.column_dimensions['B'].width = 6
        for j in range(3, NP + 4): ws.column_dimensions[L(j)].width = 8.5
        ws.freeze_panes = 'C4'; setup(ws)

# ------------------------------------------------------------------ 5. long list
ws = wb.create_sheet('All_Heads')
hdr(ws, 1, 1, ['Wellpad', 'Model', 'Condition', 'Load combination', 'Design', 'Pile', 'Y (m)', 'Z (m)', 'FX (kN)', 'FY (kN)', 'FZ (kN)',
               'MX (kNm)', 'MY (kNm)', 'MZ (kNm)', 'M res. (kNm)', 'V res. (kN)', 'Defl. (mm)'])
for i, k in enumerate(sorted(head, key=lambda k: (k[0], k[1], ORDER.index(k[2]), k[3])), 2):
    h = head[k]
    vals = [k[0], k[1], COND[k[1]], k[2], 'LRFD' if k[2].startswith('LRFD') else 'ASD', k[3], XY[k[3]][0], XY[k[3]][1]] + \
           [round(h[q], 3) for q in ('FX', 'FY', 'FZ', 'MX', 'MY', 'MZ', 'M', 'V', 'd')]
    for j, v in enumerate(vals, 1): cell(ws, i, j, v, fmt='#,##0.0' if j >= 9 else ('0.00' if j in (7, 8) else None))
for j, w in enumerate([8, 7, 13, 14, 7, 6, 7, 7, 10, 9, 9, 9, 10, 10, 10, 9, 9], 1): ws.column_dimensions[L(j)].width = w
ws.freeze_panes = 'A2'; ws.auto_filter.ref = f'A1:Q{len(head) + 1}'
wb.save(out)

if __name__ == '__main__':
    assert NP == 16 and XY[1] == (-8.75, -4.75) and XY[16] == (8.75, 4.75)
    assert abs(head[('WPA', 'M1', 'ASD2(max)', 1)]['FX'] - 5053.5) < 0.05     # = Rekap_Wellpad_A (Drive)
    for (wp, mdl), d in DATA.items():                                           # equilibrium
        for c in d['cases']:
            assert abs(sum(head[(wp, mdl, rerun.lcname(c['name']), p)]['FX'] for p in XY) - c['load'][0]) < 0.001 * c['load'][0]
    print('ok', out, len(head), 'pile-head records')
