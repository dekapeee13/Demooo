"""Isi template 'Analisis_Konsolidasi.xlsx' (bearing capacity + konsolidasi + parameter PLAXIS) untuk BH-01..BH-05 Bora Pulu.
Edit level XML (gambar, chart, validasi, format template tetap)."""
import json, math, re, zipfile, datetime, statistics as st
import openpyxl
from lxml import etree
from xlsx_cellwriter import Sheet, N as NS

TPL = 'bc/tpl.xlsx'
OLD1, OLD2 = 'Soil Parameter (BH-04)', 'BH-04 (H=2 m) - Replacement'
AREA = {'BH-01': 'Wellpad D', 'BH-02': 'Wellpad A', 'BH-03': 'Log Yard', 'BH-04': 'Wellpad D', 'BH-05': 'Wellpad A'}
GAM = {'BH-01': (17.8, 19.2), 'BH-02': (18.5, 18.4), 'BH-03': (16.5, 19.0), 'BH-04': (19.6, 20.2), 'BH-05': (18.3, 18.2)}
GWL = {'BH-01': 9, 'BH-02': 7, 'BH-03': 8, 'BH-04': 9, 'BH-05': 7}
SRC = {'BH-04': ('BH-01', 25), 'BH-05': ('BH-02', 25), 'BH-03': ('BH-02', 22.5)}   # data di bawah kedalaman bor = salinan
NLAY = 12
TODAY = (datetime.date(2026, 10, 6) - datetime.date(1899, 12, 30)).days

lab = json.load(open('lab_rows.json'))
num = lambda v: float(v) if isinstance(v, (int, float)) else None
liq = openpyxl.load_workbook('liqlat.xlsx', data_only=True)
DESC = {}
for b in AREA:
    ws = liq[b]
    DESC[b] = {round(ws.cell(r, 2).value, 2): (ws.cell(r, 3).value, str(ws.cell(r, 4).value)) for r in range(90, 104)
               if isinstance(ws.cell(r, 2).value, (int, float))}


def cat_of(desc, typ):
    d = (desc or '').strip().upper()
    if d:
        t = d.split()[0]
        if t[0] in 'SG':
            return 'Sand'
        if t.startswith('ML') or t.startswith('CL-ML') or t.startswith('MH'):
            return 'Silt'
        if t.startswith('C'):
            return 'Clay'
    return 'Sand' if typ == 'Sand' else 'Clay'


def phi_sand(n):
    n60 = min(1.083 * n, 60)
    return round(27.1 + 0.3 * n60 - 0.00054 * n60 ** 2, 1)          # Wolff (1989)


def phi_fine(pi):
    pi = max(pi or 10, 5)
    return round(min(max(math.degrees(math.asin(0.8 - 0.094 * math.log(pi))), 20), 30), 1)  # Terzaghi, Peck & Mesri (1996)


def cv_of(ll):
    pts = [(30, 3e-3), (40, 1.5e-3), (50, 7e-4), (60, 4e-4), (80, 2e-4), (100, 1.2e-4), (120, 8e-5)]   # NAVFAC DM-7.1 Fig.5
    ll = min(max(ll, 30), 120)
    for (a, ca), (b, cb) in zip(pts, pts[1:]):
        if a <= ll <= b:
            return 10 ** (math.log10(ca) + (ll - a) / (b - a) * (math.log10(cb) - math.log10(ca)))


def sublayers(bh):
    wt = openpyxl.load_workbook(f'drv2/PBCP-{bh}_rev3.xlsx')['Data Tanah']
    out = []
    for r in range(7, 48):
        top, thk, typ, n = wt[f'B{r}'].value, wt[f'D{r}'].value, wt[f'F{r}'].value, wt[f'G{r}'].value
        src = SRC[bh][0] if bh in SRC and top >= SRC[bh][1] else bh
        ln, ld = DESC[src].get(round(top, 2), (None, None))
        if ln != n:          # deskripsi laporan likuefaksi dipakai hanya bila N-nya sama dengan data SPT
            ld = None
        out.append(dict(top=float(top), bot=float(top) + float(thk), N=float(n or 0), src=src, cat=cat_of(ld, typ)))
    out[0]['N'] = out[1]['N'] if out[0]['N'] == 0 else out[0]['N']
    out[0]['cat'] = out[1]['cat'] if out[0]['N'] == out[1]['N'] else out[0]['cat']
    return out


def merge(sub, n=NLAY):
    L = [dict(s, h=s['bot'] - s['top']) for s in sub]
    while len(L) > n:
        best, bi = None, None
        for i in range(len(L) - 1):
            a, b = L[i], L[i + 1]
            c = (0 if a['cat'] == b['cat'] else 100) + abs(a['N'] - b['N']) + 0.2 * (a['h'] + b['h'])
            if best is None or c < best:
                best, bi = c, i
        a, b = L[bi], L[bi + 1]
        h = a['h'] + b['h']
        cat = a['cat'] if a['h'] >= b['h'] else b['cat']
        L[bi] = dict(top=a['top'], bot=b['bot'], h=h, N=(a['N'] * a['h'] + b['N'] * b['h']) / h, cat=cat,
                     src=a['src'] if a['h'] >= b['h'] else b['src'])
        del L[bi + 1]
    for l in L:
        l['N'] = round(l['N'])
    return L


def labstats(bh, top, bot):
    def pick(b, t0, t1):
        return [d for d in lab if d['B'] == b and t0 <= (d['C'] + d['D']) / 2 < t1]
    s = pick(bh, top, bot)
    if bh in SRC and bot > SRC[bh][1]:
        s += pick(SRC[bh][0], max(top, SRC[bh][1]), bot)
    avg = lambda k: round(st.mean([num(d.get(k)) for d in s if num(d.get(k)) is not None]), 3) if any(num(d.get(k)) is not None for d in s) else None
    return dict(Gs=avg('K'), w=avg('L'), e=avg('O'), Sr=avg('P'), LL=avg('Q'), PL=avg('R'), PI=avg('S'), n=len(s))


def fine_default(bh):
    s = [d for d in lab if d['B'] == bh and num(d.get('Q')) is not None] or [d for d in lab if num(d.get('Q')) is not None]
    m = lambda k: round(st.mean([num(d[k]) for d in s if num(d.get(k)) is not None]), 3)
    return dict(Gs=m('K'), w=m('L'), e=m('O'), Sr=m('P'), LL=m('Q'), PL=m('R'), PI=m('S'))


def unshare(xml_bytes, ws):
    """ganti shared formula dengan formula eksplisit (agar aman ditimpa)"""
    tr = etree.fromstring(xml_bytes)
    for c in tr.iter(NS + 'c'):
        f = c.find(NS + 'f')
        if f is not None and f.get('t') == 'shared':
            v = ws[c.get('r')].value
            if isinstance(v, str) and v.startswith('='):
                for k in ('t', 'ref', 'si'):
                    f.attrib.pop(k, None)
                f.text = v[1:]
        if f is not None:
            vv = c.find(NS + 'v')
            if vv is not None:
                c.remove(vv)
            if c.get('t') in ('str', 'e', 'b', 'n'):
                del c.attrib['t']
    return etree.tostring(tr, xml_declaration=True, encoding='UTF-8', standalone=True)


def build(bh):
    g_un, g_sat = GAM[bh]; gwl = GWL[bh]
    sub = merge(sublayers(bh), 41)          # sublapisan 1,5 m (tanpa merge)
    lay = merge(sublayers(bh))
    fd = fine_default(bh)
    z = zipfile.ZipFile(TPL)
    wbf = openpyxl.load_workbook(TPL)
    S1 = Sheet(unshare(z.read('xl/worksheets/sheet1.xml'), wbf.worksheets[0]))
    S2 = Sheet(unshare(z.read('xl/worksheets/sheet2.xml'), wbf.worksheets[1]))
    W1 = lambda r, v: S1.set(r, v, allow_formula=True)
    W2 = lambda r, v: S2.set(r, v, allow_formula=True)
    F1, F2 = S1.set_formula, S2.set_formula
    new2 = f'{bh} (H=2 m) - Replacement'
    log = []

    # ---------- parameter tiap lapisan ----------
    sig = 0; prev_mid = 0; P = []
    for k, l in enumerate(lay):
        ls = labstats(bh, l['top'], l['bot'])
        fine = l['cat'] != 'Sand'
        g = ls  # lab
        mid = (l['top'] + l['bot']) / 2
        gam = g_sat if mid > gwl else g_un
        sig += (mid - prev_mid) * gam; prev_mid = mid
        s_eff = sig - 10 * max(mid - gwl, 0)
        LL = (g['LL'] or fd['LL']) if fine else None
        PL = (g['PL'] or fd['PL']) if fine else None
        PI = (LL - PL) if fine else None
        e0 = g['e'] or (fd['e'] if fine else 0.65)
        cu = 5 * l['N'] if fine else None
        Cc = round(0.009 * (LL - 10), 3) if fine else 0
        Cr = round(0.2 * Cc, 3) if fine else 0
        Pc = round(cu / 0.22) if fine else None
        C = Cr if (fine and Pc > s_eff) else Cc
        mv = round(0.435 * C / ((1 + e0) * max(s_eff, 10) / 98.07), 4) if fine else None
        Cv = cv_of(LL) if fine else None
        phi = phi_fine(PI) if fine else phi_sand(l['N'])
        w = g['w'] or (fd['w'] if fine else None)
        P.append(dict(l, mid=mid, gam=gam, LL=LL, PL=PL, PI=PI, e0=round(e0, 3), cu=cu, Cc=Cc, Cr=Cr, Pc=Pc, mv=mv, Cv=Cv,
                      phi=phi, c=round(0.1 * cu, 1) if fine else 0, Gs=g['Gs'] or (fd['Gs'] if fine else 2.65), w=w,
                      Sr=g['Sr'], nlab=g['n'], s_eff=s_eff,
                      LI=round((w - PL) / PI, 3) if (fine and w and PI) else 0))

    # ================= SHEET 1: Soil Parameter =================
    fines_lab = sorted([d for d in lab if d['B'] == bh and num(d.get('Q')) is not None], key=lambda d: d['C'])
    for i in range(24):
        r = 3 + i
        if i < len(fines_lab):
            d = fines_lab[i]
            LLs = num(d['Q']); e = num(d['O']) or fd['e']
            ns = num(d.get('G')) or num(d.get('H')) or 10
            cc = round(0.009 * (LLs - 10), 3)
            dm = (d['C'] + d['D']) / 2
            sv = g_un * min(dm, gwl) + g_sat * max(dm - gwl, 0) - 10 * max(dm - gwl, 0)
            pc = 5 * ns / 0.22
            C = 0.2 * cc if pc > sv else cc
            for ref, v in (('B', bh), ('C', d['C']), ('D', '-'), ('E', d['D']), ('F', cc), ('G', round(0.2 * cc, 3)),
                           ('H', round(e, 3)), ('I', round(pc / 98.07, 2)),
                           ('J', round(0.435 * C / ((1 + e) * sv / 98.07), 4)), ('K', float(f'{cv_of(LLs):.2e}'))):
                W1(f'{ref}{r}', v)
        else:
            for ref in 'BCDEFGHIJK':
                W1(f'{ref}{r}', None)
    W1('M2', 'Catatan: SI Bora Pulu tidak memiliki uji konsolidasi (oedometer).')
    W1('M3', 'e0 = hasil uji lab (indeks). Cc = 0,009(LL-10) (Terzaghi & Peck 1967); Cr = 0,2Cc;')
    W1('M4', "Pc = cu/0,22 (Mesri 1975) dgn cu = 5N; mv = 0,435C/[(1+e0)σ'v]; Cv dari LL (NAVFAC DM-7.1).")
    W1('M5', f'Sampel lab berbutir halus {bh}: {len(fines_lab)}')
    W1('C28', gwl)
    for k in range(NLAY):
        r = 32 + k
        p = P[k]
        W1(f'C{r}', p['top'])
        W1(f'G{r}', p['N']); W1(f'H{r}', p['cat'].lower())
        W1(f'I{r}', p['gam'])
        W1(f'J{r}', p['Gs']); W1(f'K{r}', p['w']); W1(f'L{r}', p['e0'])
        W1(f'M{r}', p['Sr']); W1(f'N{r}', round(p['e0'] / (1 + p['e0']) * 100, 2))
        W1(f'O{r}', p['LL']); W1(f'P{r}', p['PL'])
        F1(f'R{r}', f'IFERROR((K{r}-P{r})/Q{r},"")')
        W1(f'T{r}', p['cu']); W1(f'U{r}', p['c']); W1(f'V{r}', p['phi'])
        W1(f'W{r}', p['Cc']); W1(f'X{r}', p['Cr'])
        F1(f'Z{r}', f'10*MAX(D{r}-$C$28,0)')
        W1(f'AB{r}', p['Pc'])
        F1(f'AC{r}', f'IFERROR(AB{r}/AA{r},"")')
        F1(f'AD{r}', f'IF(AB{r}="","-",IF(AB{r}>(AA{r}),"OC","NC-OC"))')
        W1(f'AE{r}', p['mv']); W1(f'AF{r}', float(f"{p['Cv']:.2e}") if p['Cv'] else None)
    W1('C44', P[-1]['bot'])
    F1('AF44', 'IFERROR(AVERAGE(AF32:AF43),0.001)')
    F1('AF45', 'AF44')
    F1('C45', f"'{new2}'!AA101")
    W1('B47', 'Korelasi: γ = Tabel 5.1 laporan likuefaksi (atas/bawah MAT); cu = 5N; c\' = 0,1cu; φ\' pasir = Wolff (1989), N60 = 1,083N;')
    W1('B48', "φ' lanau/lempung = sin φ' = 0,8 - 0,094 ln(PI) (Terzaghi, Peck & Mesri 1996); LL/PL/Gs/w/e0 = rata-rata uji lab pada lapisan, bila kosong rata-rata BH.")
    if bh in SRC:
        W1('B49', f'Lapisan di bawah {SRC[bh][1]} m memakai data {SRC[bh][0]} (bor {bh} hanya s/d {SRC[bh][1]} m).')

    # ================= SHEET 2: perhitungan =================
    W2('J2', 'GEOTHERMAL WELLPAD PSPE BORA PULU - SIGI, SULAWESI TENGAH')
    W2('J3', 'PT EDC PANAS BUMI INDONESIA')
    W2('J4', f'{bh} ({AREA[bh]})')
    W2('AC3', TODAY); W2('AC5', TODAY)
    W2('AA37', bh); W2('AA40', gwl)
    # Tabel 1: 5 m teratas per 1 m dari SPT
    for i in range(5):
        r = 46 + i; zz = i + 0.5
        s = next(x for x in sub if x['top'] <= zz < x['bot'])
        if s['N'] == 0:
            s = sub[sub.index(s) + 1]
        fine = s['cat'] != 'Sand'
        ls = labstats(bh, s['top'], s['bot'])
        pi = ((ls['PI'] if ls['PI'] is not None else fd['PI'])) if fine else None
        W2(f'J{r}', s['N']); W2(f'L{r}', round(pi, 1) if pi is not None else '-')
        W2(f'N{r}', s['cat']); W2(f'Q{r}', g_un if i + 1 <= gwl else g_sat)
        W2(f'W{r}', 5 * s['N'] if fine else 0)
        W2(f'AC{r}', phi_fine(pi) if fine else phi_sand(s['N']))
    # daya dukung: undrained (dominan lanau/lempung) atau drained (dominan pasir)
    F2('AA75', 'IFERROR(SUMPRODUCT((N46:N50<>"Sand")*W46:W50*D46:D50)/SUMPRODUCT((N46:N50<>"Sand")*D46:D50),0)')
    F2('AA148', 'IFERROR(SUMPRODUCT((N140:N144<>"Sand")*W140:W144*D140:D144)/SUMPRODUCT((N140:N144<>"Sand")*D140:D144),0)')
    W2('AG75', 'tebal lanau/lempung 0-5 m'); F2('AH75', 'SUMPRODUCT((N46:N50<>"Sand")*D46:D50)')
    F2('AA77', 'IF(AH75=0,SUMPRODUCT(AC46:AC50,D46:D50)/SUM(D46:D50),0)')
    F2('AA79', 'IF(AA77=0,1,EXP(PI()*TAN(RADIANS(AA77)))*TAN(RADIANS(45+AA77/2))^2)')
    F2('AA80', 'IF(AA77=0,5.14,(AA79-1)/TAN(RADIANS(AA77)))')
    F2('AA81', '2*(AA79+1)*TAN(RADIANS(AA77))')
    W2('B83', 'Lebar dasar timbunan/platform (kondisi drained)'); W2('X83', 'B ='); W2('AA83', 10); W2('AD83', 'm')
    F2('B86', 'IF(AA77=0,"Ultimate bearing capacity (Undrained Condition),","Ultimate bearing capacity (Drained Condition),")')
    F2('U86', 'IF(AA77=0,"Qu = (Cu x Nc),","Qu = c\'Nc + 0,5gBNg,")')
    F2('AA86', 'IF(AA77=0,AA75*AA80,AA76*AA80+0.5*AA82*AA83*AA81)')
    F2('B87', 'IF(AA77=0,"Allowable bearing capacity (Undrained Condition),","Allowable bearing capacity (Drained Condition),")')
    # replacement
    for i, r in enumerate(range(141, 145)):
        F2(f'N{r}', f"'{new2}'!N{47 + i}")
    W2('AG148', 'tebal lanau/lempung'); F2('AH148', 'SUMPRODUCT((N141:N144<>"Sand")*D141:D144)')
    F2('AA150', 'IF(AH148=0,SUMPRODUCT(AC140:AC144,D140:D144)/SUM(D140:D144),0)')
    F2('AA152', 'IF(AA150=0,1,EXP(PI()*TAN(RADIANS(AA150)))*TAN(RADIANS(45+AA150/2))^2)')
    F2('AA153', 'IF(AA150=0,5.14,(AA152-1)/TAN(RADIANS(AA150)))')
    F2('AA154', '2*(AA152+1)*TAN(RADIANS(AA150))')
    F2('AA155', 'IF(AA40>5,SUMPRODUCT(Q140:Q144,D140:D144)/SUM(D140:D144),SUMPRODUCT(Q140:Q144,D140:D144)/SUM(D140:D144)-10)')
    F2('U159', 'IF(AA150=0,"Qu = (Cu x Nc),","Qu = c\'Nc + 0,5gBNg,")')
    F2('AA159', 'IF(AA150=0,AA148*AA153,AA149*AA153+0.5*AA155*AA83*AA154)')
    for r in range(138, 143):
        F2(f'AJ{r}', f'VLOOKUP(AI{r},$AH$129:$AJ$132,2,FALSE)')
    # konsolidasi
    for k in range(NLAY):
        r = 185 + k; p = P[k]
        W2(f'D{r}', p['bot']); W2(f'F{r}', p['N']); W2(f'H{r}', p['cat'].upper())
        W2(f'J{r}', round(p['gam'] - (10 if p['mid'] > gwl else 0), 2))
        W2(f'L{r}', p['e0']); W2(f'N{r}', p['Cc']); W2(f'P{r}', p['Cr'])
        W2(f'R{r}', p['Pc'] or 0); W2(f'AG{r}', p['LI']); W2(f'AO{r}', p['Pc'] or 0)
        # σ'0 di tengah lapisan (kumulatif penuh) dan σ'c = max(Pc, σ'0) agar konsisten dengan kolom Pc
        F2(f'V{r}', f'(D{r}-D{r-1})/2*J{r}' if r == 185 else f'V{r-1}+(D{r-1}-D{r-2})/2*J{r-1}+(D{r}-D{r-1})/2*J{r}')
        F2(f'X{r}', f'MAX(R{r},V{r})')
    fines = [p for p in P if p['cat'] != 'Sand' and p['bot'] <= 30]
    hc = max(p['bot'] for p in fines) if fines else P[0]['bot']
    W2('AA202', hc)
    W2('X209', 'Double drainage')
    F2('AA211', f"'Soil Parameter ({bh})'!AF44")
    # PLAXIS parameter table
    W2('AA307', NLAY)
    F2('D329', 'AA90'); W2('D330', 0); F2('D331', '-AA128')
    for k in range(NLAY):
        r = 331 + k; p = P[k]
        if k > 0:
            W2(f'D{r}', -p['top'])
        n = p['N']
        W2(f'H{r}', n); W2(f'J{r}', p['cat'].lower())
        W2(f'N{r}', g_sat)
        W2(f'P{r}', 1000 * n); W2(f'R{r}', 800 * n); W2(f'T{r}', 3000 * n); W2(f'V{r}', 0.3)
        W2(f'X{r}', p['cu'] if p['cu'] else '-'); W2(f'Z{r}', p['c'] if p['cat'] != 'Sand' else 1)
        W2(f'AB{r}', p['phi'])
        W2(f'AD{r}', {'Sand': 0.0864, 'Silt': 0.00432, 'Clay': 0.0004548}[p['cat']])
    W2('AG343', -P[-1]['bot'])
    # hasil PLAXIS -> dikosongkan (belum dianalisis untuk lokasi ini)
    for ref in ('AA365', 'AA366', 'AA367', 'AA389', 'AA390', 'AA391', 'AA405', 'AA438', 'AA439', 'AA440', 'Q515', 'U516', 'U517'):
        W2(ref, None)
    for ref, cond in (('X369', 'AA367'), ('X393', 'AA391'), ('X442', 'AA440')):
        f = S2.cell(ref).find(NS + 'f').text
        F2(ref, f'IF({cond}="","Belum dianalisis (PLAXIS)",{f})')
    F2('X407', 'IF(AA405="","Belum dianalisis (PLAXIS)",IF(AA405<AA406,"T < Tall…. OK!!","T > Tall….NG!!"))')
    for r, src in ((450, (365, 366, 367)), (451, (389, 390, 391)), (452, (438, 439, 440))):
        for col, chk, s in (('N', 'Q', src[0]), ('T', 'W', src[1]), ('Z', 'AC', src[2])):
            F2(f'{col}{r}', f'IF(AA{s}="","-",AA{s})')
            F2(f'{chk}{r}', f'IF({col}{r}="-","-",IF({col}{r}>=$K${r},"OK","NOT OK"))')
    W2('AC516', '-'); F2('AC517', 'IF(U517="","-",IF((U517-Q515)<=Y517,"OK","NOT OK"))')
    F2('M506', 'AA90')
    W2('AA413', 0.64); W2('AA414', 1.1)
    W2('AF365', 'Isi dari hasil PLAXIS 2D (gambar di halaman ini masih dari template).')
    W2('AF389', 'Isi dari hasil PLAXIS 2D.'); W2('AF438', 'Isi dari hasil PLAXIS 2D.'); W2('AF515', 'Isi dari hasil PLAXIS 2D.')

    # ---------- tulis paket ----------
    edited = {'xl/worksheets/sheet1.xml': S1.bytes(), 'xl/worksheets/sheet2.xml': S2.bytes()}
    dst = f'out/Analisis_DayaDukung_Konsolidasi_{bh}.xlsx'
    rep = [(OLD2.encode(), new2.encode()), (OLD1.encode(), f'Soil Parameter ({bh})'.encode())]
    with zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED) as zo:
        for it in z.infolist():
            if it.filename == 'xl/calcChain.xml':
                continue
            data = edited.get(it.filename) or z.read(it.filename)
            if it.filename.endswith('.xml') or it.filename.endswith('.rels'):
                for a, b in rep:
                    data = data.replace(a, b)
                    data = data.replace(a.replace(b"'", b"&apos;"), b.replace(b"'", b"&apos;"))
            if it.filename == 'xl/workbook.xml':
                s = data.decode('utf8')
                s = re.sub(r'<calcPr([^>]*?)/>', lambda m: '<calcPr' + re.sub(r'\s*fullCalcOnLoad="[^"]*"', '', m.group(1)) + ' fullCalcOnLoad="1"/>', s, 1)
                data = s.encode()
            elif it.filename == 'xl/_rels/workbook.xml.rels':
                data = re.sub(rb'<Relationship [^>]*calcChain[^>]*/>', b'', data)
            elif it.filename == '[Content_Types].xml':
                data = re.sub(rb'<Override [^>]*calcChain[^>]*/>', b'', data)
            zo.writestr(it, data)
    return dst, P, hc


if __name__ == '__main__':
    import sys
    for bh in (sys.argv[1:] or AREA):
        dst, P, hc = build(bh)
        print(dst, 'hc', hc)
        for p in P:
            print(f"  {p['top']:5.1f}-{p['bot']:5.1f} {p['cat']:5s} N{p['N']:3d} g{p['gam']} e0 {p['e0']} cu {p['cu']} phi {p['phi']} Cc {p['Cc']} Pc {p['Pc']} lab{p['nlab']}")
